from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from doc_register.rent_processor import (
    RENT_COLUMNS,
    RENT_SHEET_NAME,
    _eligible_contracts,
    _is_lease_contract_row,
    _upsert_result,
    extract_rent_clause,
)


def test_extracts_only_the_two_requested_clause_five_values() -> None:
    result = extract_rent_clause(
        """[Page 8]
CLÁUSULA 5.ª
Renda e Forma de Pagamento
1. A renda anual devida pela Arrendatária ao Senhorio corresponde a 1.000,00 EUR.
9. A Arrendatária fica obrigada ao Pagamento anual de 25% do valor da renda anual
até à emissão da licença de construção, a título de reserva.
CLÁUSULA 6.ª
O valor de 99.999,00 EUR e 80% não pertencem à cláusula quinta.
"""
    )

    assert result.status == "processed"
    assert result.annual_rent.value == 1000
    assert result.annual_rent.display == "1.000,00 EUR"
    assert result.annual_rent.page == 8
    assert result.reservation_percent.value == 25
    assert result.reservation_percent.display == "25%"
    assert result.reservation_percent.page == 8


def test_does_not_extract_values_outside_clause_five() -> None:
    result = extract_rent_clause(
        """[Page 8]
CLÁUSULA 4.ª
A renda anual corresponde a 1.000,00 EUR e 25% é pago a título de reserva.
CLÁUSULA 6.ª
A renda anual corresponde a 2.000,00 EUR e 30% é pago a título de reserva.
"""
    )

    assert result.status == "needs_review_not_found"


def test_preserves_rate_per_hectare_without_inventing_a_total_or_area() -> None:
    result = extract_rent_clause(
        """[Page 8]
Clausula 5 (Renda e Forma de Pagamento)
A renda anual corresponde a 1 000,00 EUR por hectare efetivamente ocupado.
[Page 9]
Pagamento anual de 25 por cento do valor da renda anual, a título de reserva,
após aprovação ambiental e emissão da Licença de Produção, até à emissão da Licença de Construção.
O pagamento é proporcional e o excesso será compensado com a primeira renda.
Cláusula 6.ª
"""
    )

    assert result.rent_amount_type == "rate_per_area"
    assert result.rent_rate_eur_per_ha_year == 1000
    assert result.annual_rent_total_eur is None
    assert result.effective_occupied_area_ha is None
    assert result.reservation_percent.value == 25
    assert result.status == "processed"


def test_extracts_structured_reservation_events_across_pages() -> None:
    result = extract_rent_clause(
        """[Page 8]
CLÁUSULA 5.ª — Renda e Forma de Pagamento
A renda anual é de 1.000,00 EUR por hectare efetivamente ocupado.
[Page 9]
A Arrendatária pagará anualmente 25% do valor da renda anual, a título de reserva,
após aprovação ambiental e emissão da Licença de Produção, até à emissão da Licença de Construção.
CLÁUSULA 6.ª
"""
    )

    record = result.as_financial_record()
    assert record["rent_rate_eur_per_ha_year"] == 1000.0
    assert record["annual_rent_formula"] == "rent_rate_eur_per_ha_year * effective_occupied_area_ha"
    assert record["reservation_rate"] == 0.25
    assert record["reservation_start_triggers"] == ["environmental_approval", "production_license_issued"]
    assert record["reservation_end_trigger"] == "construction_license_issued"
    assert record["source_pages"] == [8, 9]


def test_excel_eligibility_accepts_only_lease_classification() -> None:
    assert _is_lease_contract_row({"document_category": "lease_contract"})
    assert _is_lease_contract_row({"document_type": "Contrato de Arrendamento Rural"})
    assert not _is_lease_contract_row({"document_type": "Contrato Promessa de Compra e Venda"})


def test_reads_preclassified_contract_rows_and_upserts_result_sheet() -> None:
    with TemporaryDirectory() as directory:
        base = Path(directory)
        pdf = base / "contract.pdf"
        pdf.write_bytes(b"placeholder")
        workbook_path = base / "register.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Document Register"
        sheet.append(["source_file_name", "copied_file_path", "sha256", "document_type"])
        sheet.append([pdf.name, str(pdf), "digest-a", "Contrato de Arrendamento"])
        sheet.append(["other.pdf", str(base / "other.pdf"), "digest-b", "Fatura"])
        workbook.save(workbook_path)
        workbook.close()

        contracts = _eligible_contracts(workbook_path, base)
        assert len(contracts) == 1
        assert contracts[0].source_file_path == pdf

        payload = {column: "" for column in RENT_COLUMNS}
        payload.update({"sha256": "digest-a", "status": "processed", "annual_rent_eur": 1000.0})
        _upsert_result(workbook_path, payload)
        payload["annual_rent_eur"] = 1200.0
        _upsert_result(workbook_path, payload)

        output = load_workbook(workbook_path, data_only=True)
        result_sheet = output[RENT_SHEET_NAME]
        assert result_sheet.max_row == 2
        assert result_sheet.cell(row=2, column=RENT_COLUMNS.index("annual_rent_eur") + 1).value == 1200.0
        output.close()
