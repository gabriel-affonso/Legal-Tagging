"""Semantic regressions from the 5.1 implementation prompt; no live model required."""
import json
import time
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from doc_register.step55.canonical import Assertion, CanonicalResult, ExtractionTask
from doc_register.step55.document_map import Observation
from doc_register.step55.pipeline import Pipeline, PipelineConfig
from doc_register.step55.models import LocalModels, validate_response
from doc_register.step55.resolution import reconcile
from doc_register.step55.review import apply_decisions
from doc_register.step55.export import export_staging
from doc_register.step55.values import validate_value, calendar_offset, evaluate_formula
from doc_register.step55.evaluation import evaluate
from test_step55 import make_pdf

PARCELS = """Prédio denominado Fonte da Carvalha, artigo matricial: 68, descrição predial: 529,
freguesia de Castedo, concelho de Torre de Moncorvo, área total de 25.650 m2.
Prédio denominado Cascalheira, artigo matricial: 306, descrição predial: 530,
freguesia de Castedo, concelho de Torre de Moncorvo, área total de 11.502 m2.
Área útil aproximada de 3,7 ha."""
RECEIPT = """Envio de recibo de renda
18 de Abril de 2024
Valor bruto de 3.700,00 EUR
Retido na fonte de 925,00 EUR
Valor líquido de 2.775,00 EUR
artigos 68 e 530"""


def build(tmp_path, *texts, config=None):
    obs=[Observation('r'+str(i),i+1,text,None,'native') for i,text in enumerate(texts)]
    pages=[{'number':o.page,'status':'processed','issues':[],'selected_region_ids':[o.id]} for o in obs]
    return Pipeline(tmp_path,config or PipelineConfig(ocr_enabled=False)).analyze(
        obs,pages,'a'*64,str(tmp_path/'original.pdf'),tmp_path/'run',time.monotonic()+600)


def values(result, predicate):
    return [a.typed_value for a in result.assertions if a.predicate==predicate and a.acceptance_status!='rejected']


def property_objects(result):
    return {next(a.typed_value for a in result.assertions if a.subject_id==e.id and a.predicate=='cadastral_article'):
        {a.predicate:a.typed_value for a in result.assertions if a.subject_id==e.id}
        for e in result.entities if e.kind=='property'}


@pytest.mark.parametrize('reverse', [False,True])
def test_complete_objects_never_mix_and_joint_area_is_single(tmp_path,reverse):
    text=PARCELS
    if reverse:
        parts=text.split('Prédio ')
        text='Prédio '+parts[2].split('Área útil')[0]+'Prédio '+parts[1]+'Área útil aproximada de 3,7 ha.'
    result=build(tmp_path,'CONTRATO DE ARRENDAMENTO\n'+text)
    objects=property_objects(result)
    assert set(objects)=={'68','306'}
    for article,name,description,area in [('68','Fonte da Carvalha','529','25650'),('306','Cascalheira','530','11502')]:
        p=objects[article]
        assert (p['property_name'],p['registry_description'],p['area_measurement']['amount'])==(name,description,area)
        assert p['property_parish']=='Castedo'
        assert p['property_municipality']=='Torre de Moncorvo'
    useful=[a for a in result.assertions if a.predicate=='area_measurement' and a.typed_value['concept']=='leased_usable']
    assert len(useful)==1
    assert next(e.kind for e in result.entities if e.id==useful[0].subject_id)=='lease_object'
    assert useful[0].typed_value['approximate'] is True
    assert next(d for d in result.derivations if d.operation=='sum_contract_stated_property_areas').value['amount']=='3.7152'


def test_property_continues_across_pages(tmp_path):
    result=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nPrédio denominado Fonte da Carvalha,',
        'artigo matricial: 68, descrição predial: 529, área total de 25.650 m2.')
    assert len(result.logical_documents)==1
    obj=property_objects(result)['68']
    assert obj['property_name']=='Fonte da Carvalha'
    area=next(a for a in result.assertions if a.predicate=='area_measurement')
    evidence={e.id:e for e in result.evidence}
    assert {evidence[i].page for i in area.evidence_ids}=={1,2}


def test_document_mentions_do_not_reclassify_contract(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nO titular do cartão de cidadão assinou.\nO prédio está descrito na conservatória. A planta está anexa.',
        'O contrato terá a duração de 29 anos e 11 meses.')
    assert len(r.logical_documents)==1
    assert values(r,'term_duration')[0]['years']==29


def test_rgpd_boundary_and_legal_article_not_property(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nArtigo 306 do Código Civil.\nDeclaração de proteção de dados\nMiraflores Office Center, IBAN/NIB, artigo 68.')
    assert [d.source_type for d in r.logical_documents]==['lease_contract','privacy_notice']
    assert not values(r,'cadastral_article') and not values(r,'iban')


def test_namespace_mismatch_and_companies_never_merged(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nArrendatária: GESTO ENERGIA, S.A.\n'+PARCELS,
        RECEIPT+'\nPagador: Malhada Green, S.A.')
    assert len([e for e in r.entities if e.kind=='organization'])==2
    ref=next(a for a in r.assertions if a.predicate=='related_property_article' and a.typed_value=='530')
    assert 'reference_namespace_or_value_mismatch' in ref.validation_issues
    assert not [r for r in r.relations if r.predicate=='same_as']
    assert values(r,'document_date')[0]['value']=='2024-04-18'
    assert not values(r,'bank_execution_date')
    assert values(r,'payment_status')==['reported']


def test_operational_and_cadastral_area_never_overwrite(tmp_path):
    r=build(tmp_path,'CADERNETA PREDIAL RÚSTICA\nARTIGO MATRICIAL: 68\nÁrea total: 25.650 m2')
    original=next(a for a in r.assertions if a.predicate=='area_measurement')
    other=original.model_copy(deep=True)
    other.id='operational';other.origin_kind='operacional'
    other.typed_value={'amount':'2.55','unit':'ha','concept':'operational_contracted','approximate':False}
    r.assertions.append(other);reconcile(r)
    assert len(values(r,'area_measurement'))==2
    assert all(a.resolution_status!='conflict' for a in r.assertions)


def test_roles_no_generic_entities_or_conjugal_inference(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nSenhorio: Parte Incumpridora\nArrendatario: GESTO ENERGIA, S.A.\nMaria Ribeiro, NIF 196415209, casada com Manuel Ribeiro.\nrepresentada por Pedro Miguel Fernandes, na qualidade de administrador.')
    assert 'Parte Incumpridora' not in [e.label for e in r.entities]
    assert values(r,'contract_role')==['lessee']
    assert not [a for a in r.assertions if a.predicate=='contract_role' and next(e.label for e in r.entities if e.id==a.subject_id).startswith(('Manuel','Pedro'))]


def test_temporal_references_remain_events_and_deadlines(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nAssinado em 22 de Dezembro de 2020\nO contrato terá a duração de 29 anos e 11 meses.\nA Condição Suspensiva será verificada no prazo de três anos da assinatura.\nDisponibilizar no prazo de três meses após notificação.\nO início do arrendamento depende da Condição Suspensiva.\nData da sua assinatura.')
    assert len(values(r,'term_duration'))==1
    assert len(values(r,'temporal_rule'))==2
    assert values(r,'signed_date')[0]['value']=='2020-12-22'
    assert not values(r,'contract_end_date')
    assert all(d.status=='blocked' for d in r.derivations if d.operation in {'calendar_deadline','contract_end_date'})
    assert not [e for e in r.entities if e.kind=='event_occurrence']
    assert calendar_offset('2024-01-31',{'months':1}) is None
    assert calendar_offset('2024-01-31',{'months':1},convention='clamp_month_end')=='2024-02-29'


@pytest.mark.parametrize('predicate,value',[
    ('signed_date',{'kind':'duration','years':29,'months':11}),
    ('signed_date',{'value':'2024-02-30'}),('signed_date',{'value':'data da assinatura'}),
    ('payment_gross_amount',{'amount':'25%','currency':'EUR'}),
    ('tax_id','196415208'),('iban','IBAN/NIB'),('nonsense','anything'),
    ('payment_gross_amount',{'amount':'NaN'}),('payment_gross_amount',{'amount':'1e9'}),
])
def test_closed_value_schema_rejects_invalid_types(predicate,value):
    with pytest.raises(ValueError):validate_value(predicate,value)


def test_financial_rate_percentage_and_duplicate_payments(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nRenda anual de 1.000,00 EUR por hectare.\nPagamento de reserva de 25% da renda anual até à licença de construção.\nIndemnização de 25% da renda anual.',
        RECEIPT+'\nReferência: ABC123', 'Recibo de renda\nReferência: ABC123\nValor líquido de 2.775,00 EUR')
    assert values(r,'rent_term')[0]['per_area_unit']=='ha'
    assert values(r,'reservation_obligation')[0]['fraction']=='0.25'
    assert values(r,'percentage_reference')==[{'fraction':'0.25'}]
    assert len([e for e in r.entities if e.kind=='payment'])==2
    assert any(x.reason=='possible_duplicate_payment_do_not_sum' for x in r.relations)
    assert not any('sum_payment' in d.operation for d in r.derivations)
    assert next(d for d in r.derivations if d.operation=='rent_amount_for_explicit_basis').status=='blocked'
    inputs={'rate':{'amount':'1000','unit':'EUR/ha/year'},'area':{'amount':'3.7','unit':'ha'}}
    assert str(evaluate_formula({'op':'multiply','left':{'ref':'rate'},'right':{'ref':'area'}},inputs)[0])=='3700.0'
    with pytest.raises(ValueError):evaluate_formula({'ref':'a'},{'a':{'expression':{'ref':'a'}}})


def test_caderneta_standard_labels_and_multiple_owners(tmp_path):
    r=build(tmp_path,'CADERNETA PREDIAL RÚSTICA\nDISTRITO: 10 - LEIRIA CONCELHO: 16 - PORTO DE MOS FREGUESIA: 06 - JUNCAL\nSECÇÃO: 018  ARTIGO MATRICIAL Nº: 344\nNOME/LOCALIZAÇÃO PRÉDIO\nCASTANHAL\nÁrea Total (ha): 0,100000\nTITULARES\nIdentificação fiscal: 119292459 Nome: PEDRO DO NASCIMENTO HORTA\nTipo de titular: Propriedade plena Parte: 1/2\nIdentificação fiscal: 120287820 Nome: JOSE VIRGILIO VIEIRA\nTipo de titular: Propriedade plena Parte: 1/2')
    p=property_objects(r)['344']
    assert p['property_municipality']=='PORTO DE MOS'
    assert p['property_parish']=='JUNCAL'
    assert p['cadastral_section']=='018'
    assert p['area_measurement']['amount']=='0.1'
    owners=[a for a in r.assertions if a.predicate=='ownership']
    assert len(owners)==2
    assert all(a.resolution_status!='conflict' for a in owners)


def test_zero_candidate_caderneta_still_has_task_context(tmp_path):
    r=build(tmp_path,'CADERNETA PREDIAL RÚSTICA\nDocumento mal reconhecido, identificação ilegível.')
    assert r.tasks and r.tasks[0].evidence_ids
    assert r.tasks[0].status=='deferred'


def task_for(r):
    contract=next(e for e in r.entities if e.kind=='contract')
    ev=next(e for e in r.evidence if e.start==0 and e.end==len(e.text))
    return ExtractionTask(id='task',task_type='temporal',logical_document_id=contract.logical_document_id,
        evidence_ids=[ev.id],entity_ids=[contract.id],permitted_predicates=['signed_date'],reason='test')


@pytest.mark.parametrize('failure',['timeout','invalid_json','truncated','unavailable','oversized'])
def test_llm_failures_are_bounded_and_preserve_rules(tmp_path,failure):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\n'+PARCELS)
    task=task_for(r);before=[a.model_dump() for a in r.assertions]
    config=PipelineConfig(text_enabled=True)
    if failure=='oversized':config=PipelineConfig(text_enabled=True,context_tokens=20)
    runner=LocalModels(config)
    response={'message':{'content':'not json'},'done_reason':'length' if failure=='truncated' else 'stop'}
    side=TimeoutError('timeout') if failure=='timeout' else ConnectionRefusedError() if failure=='unavailable' else None
    with patch.object(runner,'chat',return_value=response,side_effect=side):
        runner.run(task,r,tmp_path/'models',time.monotonic()+10)
    assert task.status=='blocked' if failure=='oversized' else task.status=='technical_error'
    assert runner.calls<=2
    assert [a.model_dump() for a in r.assertions]==before
    if failure=='timeout':
        assert runner.stopped and runner.calls==1
        second=task.model_copy(update={'status':'planned'})
        runner.run(second,r,tmp_path/'models',time.monotonic()+10)
        assert second.status=='blocked'


def test_model_rejects_unknown_ids_and_bad_offsets(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nAssinado em 22 de Dezembro de 2020')
    task=task_for(r);ev=next(e for e in r.evidence if e.id==task.evidence_ids[0])
    item={'subject_id':task.entity_ids[0],'predicate':'signed_date','value':{'value':'2020-12-22'},'evidence_id':ev.id,'quote':ev.quote,'start':0,'end':len(ev.quote)}
    answer={'assertions':[item],'entity_proposals':[],'unresolved':[]}
    assert len(validate_response(answer,task,r)[0])==1
    item['start']=1
    with pytest.raises(ValueError):validate_response(answer,task,r)
    item['start']=0;item['subject_id']='invented'
    with pytest.raises(ValueError):validate_response(answer,task,r)


def test_failed_ocr_inventory_and_cache_invalidation(tmp_path):
    source=tmp_path/'empty.pdf';make_pdf(source,[''])
    pipeline=Pipeline(tmp_path/'runs')
    with patch('doc_register.step55.pipeline.ocr',side_effect=TimeoutError('simulated')):
        path=pipeline.process(source)
    r=CanonicalResult.model_validate_json(path.read_text())
    assert r.pages[0]['status']=='technical_error' and r.coverage['page:1']=='technical_error'
    assert not r.assertions and not r.observations
    with patch('doc_register.step55.pipeline.ocr',return_value=RECEIPT):
        retry=pipeline.process(source)
    assert retry!=path and path.exists()
    assert pipeline.process(source)==retry
    forced=pipeline.process(source,force=True)
    assert forced!=retry and retry.exists()
    with patch('doc_register.step55.pipeline.RULE_VERSION','changed'):
        assert pipeline.process(source)!=forced


def test_review_idempotence_conflict_correction_and_export_literal(tmp_path):
    r=build(tmp_path,RECEIPT+'\nReferência: =1+1')
    path=tmp_path/'result.json';path.write_text(r.model_dump_json())
    reference=next(a for a in r.assertions if a.predicate=='payment_reference')
    payload={'run_id':r.run_id,'source_sha256':r.source_sha256,'decisions':[{'action':'accept','target_id':reference.id,'reason':'Conferido com original','evidence_ids':reference.evidence_ids}]}
    decisions=tmp_path/'decisions.json';decisions.write_text(json.dumps(payload))
    reviewed=apply_decisions(path,decisions,'Test reviewer')
    assert apply_decisions(path,decisions,'Test reviewer')==reviewed
    output=export_staging([reviewed],tmp_path/'staging.xlsx')
    with pytest.raises(FileExistsError):export_staging([reviewed],output)
    book=load_workbook(output)
    assert book['stg.Facts'].max_row==2
    assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
    assert CanonicalResult.model_validate_json(path.read_text()).decisions==[]
    book.close()


def test_evaluation_identity_scope_evidence_and_split_leakage(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\n'+PARCELS)
    manifest={'documents':[{'id':'sample','sha256':r.source_sha256,'group_id':'family','split':'test','annotators':['reviewer'],
        'fully_annotated_fields':['registry_description'],'entities':[{'gold_id':'gold1','kind':'property','identity':{'cadastral_article':'68','property_parish':'Castedo'}}],
        'facts':[{'subject_id':'gold1','predicate':'registry_description','value':'529','evidence_pages':[1]}]}]}
    report=evaluate(manifest,{'sample':r})
    assert report['extracted']['correct']==1 and report['extracted']['false_positive']==1
    assert report['automatic']['omissions']==1
    manifest['documents'].append({**manifest['documents'][0],'id':'duplicate','split':'development'})
    with pytest.raises(ValueError):evaluate(manifest,{'sample':r})


def test_short_layout_blocks_stay_in_caderneta_and_do_not_eat_next_page(tmp_path):
    texts=['CADERNETA PREDIAL RÚSTICA','ARTIGO MATRICIAL: 344','NOME/LOCALIZAÇÃO PRÉDIO','CASTANHAL',
           'Área Total (ha): 0,1','TITULARES','Identificação fiscal: 119292459 Nome: PEDRO DO NASCIMENTO HORTA','Propriedade plena Parte: 1/1']
    observations=[Observation(str(i),1,t,None,'native') for i,t in enumerate(texts)]
    observations.append(Observation('other',2,'Comunicado municipal sobre assuntos distintos.',None,'native'))
    pages=[{'number':p,'status':'processed','issues':[],'selected_region_ids':[o.id for o in observations if o.page==p]} for p in (1,2)]
    r=Pipeline(tmp_path).analyze(observations,pages,'a'*64,str(tmp_path/'original.pdf'),tmp_path/'run',time.monotonic()+60)
    assert r.logical_documents[0].page_numbers==[1]
    assert values(r,'property_name')==['CASTANHAL']
    assert len(values(r,'ownership'))==1
    assert len([e for e in r.entities if e.kind=='property'])==1


def test_conflict_can_be_rejected_without_stale_resolution(tmp_path):
    r=build(tmp_path,RECEIPT)
    original=next(a for a in r.assertions if a.predicate=='payment_gross_amount')
    wrong=original.model_copy(deep=True);wrong.id='wrong';wrong.typed_value={'amount':'999','currency':'EUR'}
    r.assertions.append(wrong);reconcile(r)
    assert original.resolution_status=='conflict'
    path=tmp_path/'result.json';path.write_text(r.model_dump_json())
    payload={'run_id':r.run_id,'source_sha256':r.source_sha256,'decisions':[
        {'target_id':wrong.id,'action':'reject','reason':'wrong reading','evidence_ids':wrong.evidence_ids},
        {'target_id':original.id,'action':'accept','reason':'checked original','evidence_ids':original.evidence_ids}]}
    decisions=tmp_path/'decisions.json';decisions.write_text(json.dumps(payload))
    reviewed=CanonicalResult.model_validate_json(apply_decisions(path,decisions,'reviewer').read_text())
    assert next(a for a in reviewed.assertions if a.id==original.id).acceptance_status=='accepted_human'
    assert next(d for d in reviewed.derivations if d.operation=='gross_minus_withholding_equals_net').value is True


def test_successful_model_response_is_anchored_and_not_automatically_accepted(tmp_path):
    r=build(tmp_path,'CONTRATO DE ARRENDAMENTO\nAssinado em 22 de Dezembro de 2020')
    task=task_for(r);ev=next(e for e in r.evidence if e.id==task.evidence_ids[0])
    quote='22 de Dezembro de 2020';start=ev.quote.index(quote)
    item={'subject_id':task.entity_ids[0],'predicate':'signed_date','value':{'value':'2020-12-22'},'evidence_id':ev.id,'quote':quote,'start':start,'end':start+len(quote)}
    answer={'assertions':[item],'entity_proposals':[],'unresolved':[]}
    runner=LocalModels(PipelineConfig(text_enabled=True))
    with patch.object(runner,'chat',return_value={'message':{'content':json.dumps(answer)}}):
        runner.run(task,r,tmp_path/'models',time.monotonic()+60)
    assert task.status=='completed'
    a=r.assertions[-1]
    assert a.acceptance_status=='needs_review' and a.raw_value==quote
    assert a.evidence_offsets[ev.id]==(start,start+len(quote))
    item['value']={'value':'2021-12-22'}
    with pytest.raises(ValueError,match='date_not_in_model_quote'):validate_response(answer,task,r)


def test_dossier_proposes_cross_file_link_without_merging(tmp_path):
    from doc_register.step55.dossier import link_results
    left=tmp_path/'contract.pdf';right=tmp_path/'record.pdf'
    make_pdf(left,['CONTRATO DE ARRENDAMENTO\n'+PARCELS])
    make_pdf(right,['CADERNETA PREDIAL RÚSTICA\nARTIGO MATRICIAL: 68\nFreguesia: Castedo\nConcelho: Torre de Moncorvo\nÁrea total: 25.650 m2'])
    pipeline=Pipeline(tmp_path/'runs',PipelineConfig(ocr_enabled=False))
    paths=[pipeline.process(left),pipeline.process(right)]
    before=[p.read_bytes() for p in paths]
    linked=link_results(paths,tmp_path/'dossier.json')
    payload=json.loads(linked.read_text())
    assert len(payload['links'])==1 and payload['links'][0]['predicate']=='possible_same_as'
    assert payload['payment_aggregation']=='disabled'
    assert before==[p.read_bytes() for p in paths]


def test_replay_uses_observations_and_does_not_import_old_candidates(tmp_path):
    source=tmp_path/'source.pdf';make_pdf(source,[RECEIPT])
    pipeline=Pipeline(tmp_path/'runs',PipelineConfig(ocr_enabled=False))
    first=pipeline.process(source)
    old=json.loads(first.read_text());old['assertions']=[]
    legacy=tmp_path/'legacy.json';legacy.write_text(json.dumps(old))
    replayed=CanonicalResult.model_validate_json(pipeline.replay(legacy).read_text())
    assert values(replayed,'payment_gross_amount')==[{'amount':'3700','currency':'EUR'}]
    assert replayed.provenance['migration']=='observations_only'
