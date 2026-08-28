from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from doc_register.models import ExtractionResult, PdfCandidate
from doc_register.property_pipeline import (
    PropertyExtractionPipeline,
    detect_lease_contract,
    select_clause_a,
    select_considering_context,
)
from doc_register.property_intelligence import (
    CadernetaValues,
    discover_caderneta_pages,
    parse_caderneta,
    recover_from_annexes,
)
from doc_register.property_pack import (
    PropertyDocumentDescriptor,
    PropertyPackDiscovery,
    PropertyPackMatch,
)
from doc_register.property_processor import _reconcile_property_evidence
from doc_register.registry import ExcelRegister, PROPERTY_SHEET_NAME, PropertyExcelRegister


def load_tests(
    loader: unittest.TestLoader,
    tests: unittest.TestSuite,
    pattern: str,
) -> unittest.TestSuite:
    functions = (
        test_skips_documents_that_do_not_reach_lease_threshold,
        test_fast_lease_detection_uses_weighted_indicators,
        test_extracts_property_fields_from_considering_clause_a_without_llm,
        test_llm_only_receives_selected_clause_and_cannot_override_regex,
        test_invalid_article_and_section_have_zero_confidence,
        test_annex_failure_preserves_invalid_field_audit_and_confidence,
        test_invalid_property_name_is_rejected_and_triggers_recovery,
        test_context_and_clause_selectors_apply_the_expected_boundaries,
        test_property_catalogue_uses_cached_ocr_for_caderneta,
        test_sp_sequence_chain_matches_unique_crp_and_caderneta,
        test_sp_sequence_chain_uses_crp_filename_when_ocr_is_weak,
        test_same_pdf_caderneta_is_reconciled_with_contract,
        test_title_and_model_b_identify_same_pdf_caderneta_near_document_end,
        test_property_pack_matches_separate_caderneta_by_identifier,
        test_property_pack_rejects_ambiguous_candidates,
        test_discovers_and_parses_caderneta_in_document_tail,
        test_reconciles_contract_and_caderneta_with_evidence,
        test_property_output_uses_a_dedicated_worksheet_in_the_main_workbook,
    )
    return unittest.TestSuite(unittest.FunctionTestCase(function) for function in functions)


def test_skips_documents_that_do_not_reach_lease_threshold() -> None:
    result = PropertyExtractionPipeline().extract(
        "Caderneta predial: artigo 123, secção K, área de 45.230 m2."
    )

    assert result.status == "skipped"
    assert result.reason == "not_lease_contract"
    assert result.used_llm is False


def test_fast_lease_detection_uses_weighted_indicators() -> None:
    detected = detect_lease_contract(
        "O Senhorio entrega ao Arrendatário o imóvel mediante renda anual."
    )

    assert detected.score == 60
    assert detected.is_lease_contract is True


def test_extracts_property_fields_from_considering_clause_a_without_llm() -> None:
    text = """
    CONTRATO DE ARRENDAMENTO RURAL
    Entre o Senhorio e o Arrendatário, mediante renda anual.
    Considerando que:
    a) O prédio rústico denominado por Quinta da Ribeira, composto por cultura
    arvense, encontra-se inscrito na matriz sob o artigo 123, secção K, com área
    de 45.230 m².
    b) O imóvel será entregue livre de pessoas e bens.
    """

    result = PropertyExtractionPipeline().extract(text)

    assert result.status == "processed"
    assert result.property_name == "Quinta da Ribeira"
    assert result.matrix_article == "123"
    assert result.matrix_section == "K"
    assert result.area_m2 == 45230
    assert result.confidence == 100
    assert result.source_section == "Considerando que"
    assert result.source_clause == "a)"
    assert result.used_llm is False


def test_llm_only_receives_selected_clause_and_cannot_override_regex() -> None:
    seen: list[str] = []

    def recover(clause: str) -> dict[str, str]:
        seen.append(clause)
        return {
            "property_name": "Nome Inventado",
            "matrix_article": "999",
            "matrix_section": "K",
            "area_m2": "100",
        }

    text = """
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio rústico denominado por Quinta da Ribeira, composto por olival,
    encontra-se inscrito na matriz sob o artigo 123, secção K.
    b) Cláusula seguinte.
    """
    result = PropertyExtractionPipeline(recover).extract(text)

    assert result.used_llm is True
    assert "a)" in seen[0]
    assert "b)" not in seen[0]
    assert result.property_name == "Quinta da Ribeira"
    assert result.matrix_article == "123"
    assert result.matrix_section == "K"
    assert result.area_m2 is None
    assert result.confidence == 80


def test_invalid_article_and_section_have_zero_confidence() -> None:
    result = PropertyExtractionPipeline().extract("""
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio encontra-se inscrito na matriz sob o artigo ES, secção LX.
    b) Cláusula seguinte.
    """, allow_llm=False)

    assert result.matrix_article == ""
    assert result.matrix_section == ""
    assert result.confidence == 0
    assert "invalid_matrix_article" in result.audit["validation_results"]
    assert "invalid_matrix_section" in result.audit["validation_results"]


def test_invalid_property_name_is_rejected_and_triggers_recovery() -> None:
    result = PropertyExtractionPipeline().extract("""
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio denominado por (doravante designado apenas como Prédio), composto
    por olival, encontra-se inscrito na matriz sob o artigo 123, secção K.
    b) Cláusula seguinte.
    """, allow_llm=False)

    assert result.property_name == ""
    assert "invalid_property_name" in result.audit["validation_results"]


def test_annex_failure_preserves_invalid_field_audit_and_confidence() -> None:
    contract = PropertyExtractionPipeline().extract("""
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio encontra-se inscrito na matriz sob o artigo ES, secção LX.
    b) Cláusula seguinte.
    """, allow_llm=False)

    result = recover_from_annexes(contract, "")

    assert result.confidence == 0
    assert "invalid_matrix_article" in result.audit["validation_results"]
    assert "invalid_matrix_section" in result.audit["validation_results"]
    assert "caderneta_not_found" in result.audit["validation_results"]


def test_context_and_clause_selectors_apply_the_expected_boundaries() -> None:
    text = "preâmbulo\nConsiderando que:\na) alvo\nb) fora do alvo"
    context, found_context = select_considering_context(text)
    clause, found_clause = select_clause_a(context)

    assert found_context is True
    assert found_clause is True
    assert clause == "a) alvo"


def test_property_pack_matches_separate_caderneta_by_identifier() -> None:
    contract = PropertyExtractionPipeline().extract("""
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio denominado por Quinta da Ribeira, composto por olival.
    b) Cláusula seguinte.
    """, allow_llm=False)
    caderneta = CadernetaValues(
        property_name="Quinta da Ribeira",
        matrix_article="269F",
        matrix_section="K",
        area_m2=14812,
    )
    discovery = PropertyPackDiscovery([])
    discovery._descriptors = (PropertyDocumentDescriptor(
        path=Path("CadernetaPredial.pdf"),
        identifiers=frozenset({"269F"}),
        kind="caderneta_predial",
        caderneta_values=caderneta,
    ),)

    match = discovery.find_for_contract(Path("VA553_269F_CA.pdf"), "", contract)

    assert match.status == "property_pack_matched"
    assert match.score >= 100
    assert match.source_path == Path("CadernetaPredial.pdf")
    assert match.caderneta_values == caderneta


def test_property_pack_rejects_ambiguous_candidates() -> None:
    caderneta = CadernetaValues(matrix_article="269F").validated()
    discovery = PropertyPackDiscovery([])
    discovery._descriptors = tuple(
        PropertyDocumentDescriptor(
            path=Path(f"Caderneta-{index}.pdf"),
            identifiers=frozenset({"269F"}),
            kind="caderneta_predial",
            caderneta_values=caderneta,
        )
        for index in range(2)
    )
    contract = PropertyExtractionPipeline().extract(
        "Contrato de arrendamento entre Senhorio e Arrendatário, com renda.",
        allow_llm=False,
    )

    match = discovery.find_for_contract(Path("VA553_269F_CA.pdf"), "", contract)

    assert match.status == "property_pack_ambiguous"


def test_property_catalogue_uses_cached_ocr_for_caderneta() -> None:
    caderneta_path = Path("CadernetaPredial-040812-R-140-J_SP18490.pdf")
    caderneta_text = """
    [Page 1] CADERNETA PREDIAL RÚSTICA
    IDENTIFICAÇÃO DO PRÉDIO
    NOME/LOCALIZAÇÃO PRÉDIO: Quinta da Ribeira
    ARTIGO MATRICIAL Nº: 269F
    SECÇÃO: K
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 1,481200
    TITULARES
    """
    discovery = PropertyPackDiscovery(
        [caderneta_path],
        text_loader=lambda path: (caderneta_text, "cached_ocr_pdf_layout_text"),
    )
    contract = PropertyExtractionPipeline().extract(
        "Contrato de arrendamento entre Senhorio e Arrendatário, com renda.",
        allow_llm=False,
    )

    match = discovery.find_for_contract(Path("VA553_269F_CA.pdf"), "", contract)

    assert match.status == "property_pack_matched"
    assert match.candidate_count == 1
    assert match.catalog_status == "caderneta_catalogued"
    assert match.catalog_text_source == "cached_ocr_pdf_layout_text"


def test_sp_sequence_chain_matches_unique_crp_and_caderneta() -> None:
    caderneta = CadernetaValues(matrix_article="140J", matrix_section="K").validated()
    crp = CadernetaValues(matrix_article="140J", matrix_section="K").validated()
    discovery = PropertyPackDiscovery([])
    discovery._descriptors = (
        PropertyDocumentDescriptor(
            path=Path("CadernetaPredial_SP18490.pdf"),
            identifiers=frozenset({"140J"}),
            kind="caderneta_predial",
            caderneta_values=caderneta,
            sharepoint_sequence=18490,
        ),
        PropertyDocumentDescriptor(
            path=Path("CRP_PR_140J_SP18491.pdf"),
            identifiers=frozenset({"PR140J", "140J"}),
            kind="crp",
            crp_values=crp,
            sharepoint_sequence=18491,
        ),
    )
    contract = PropertyExtractionPipeline().extract(
        "Contrato de arrendamento entre Senhorio e Arrendatário, com renda.",
        allow_llm=False,
    )

    match = discovery.find_for_contract(Path("PR098_CA_SP18492.pdf"), "", contract)

    assert match.status == "property_pack_matched"
    assert match.methods[0] == "sp_sequence_chain"
    assert match.crp_source_path == Path("CRP_PR_140J_SP18491.pdf")


def test_sp_sequence_chain_uses_crp_filename_when_ocr_is_weak() -> None:
    caderneta_path = Path("CadernetaPredial-040812-R-140-J_SP18490.pdf")
    crp_path = Path("CRP_PR_140J_SP18491.pdf")
    caderneta_text = """
    CADERNETA PREDIAL RÚSTICA
    IDENTIFICAÇÃO DO PRÉDIO
    NOME/LOCALIZAÇÃO PRÉDIO: Quinta da Ribeira
    ARTIGO MATRICIAL Nº: 140-J
    SECÇÃO: K
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 1,000000
    TITULARES
    """
    discovery = PropertyPackDiscovery(
        [caderneta_path, crp_path],
        text_loader=lambda path: (
            caderneta_text if path == caderneta_path else "OCR muito curto e sem marcadores úteis.",
            "cached_ocr_pdf_layout_text",
        ),
    )
    contract = PropertyExtractionPipeline().extract(
        "Contrato de arrendamento entre Senhorio e Arrendatário, com renda.",
        allow_llm=False,
    )

    match = discovery.find_for_contract(Path("PR098_CA_SP18492.pdf"), "", contract)

    assert match.status == "property_pack_matched"
    assert match.methods[0] == "sp_sequence_chain"
    assert match.caderneta_values.matrix_article == "140-J"


def test_same_pdf_caderneta_is_reconciled_with_contract() -> None:
    contract = PropertyExtractionPipeline().extract("""
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio rústico encontra-se inscrito na matriz sob o artigo 123, secção K.
    b) Cláusula seguinte.
    """, allow_llm=False)
    annex = """
    [Page 11] CADERNETA PREDIAL RÚSTICA
    IDENTIFICAÇÃO DO PRÉDIO
    NOME/LOCALIZAÇÃO PRÉDIO: Quinta da Ribeira
    ARTIGO MATRICIAL Nº: 123
    SECÇÃO: K
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 1,481200
    TITULARES
    """

    result = _reconcile_property_evidence(
        contract,
        PropertyPackMatch(status="property_pack_not_found"),
        discover_caderneta_pages(annex),
        "Contrato.pdf",
    )

    assert result.property_name == "Quinta da Ribeira"
    assert result.area_m2 == 14812
    assert result.audit["caderneta_same_pdf_found"] is True
    assert result.audit["caderneta_evidence_sources"] == ["Contrato.pdf"]


def test_title_and_model_b_identify_same_pdf_caderneta_near_document_end() -> None:
    early_contract_pages = "\n".join(
        f"[Page {index}] cláusulas do contrato" for index in range(1, 19)
    )
    caderneta = """
    [Page 19] ACTUALIZAÇÃO DE CADERNETA PREDIAL RÚSTICA
    MODELO B
    [Page 20] IDENTIFICAÇÃO DO PRÉDIO
    NOME/LOCALIZAÇÃO PRÉDIO: Herdade da Fonte
    ARTIGO MATRICIAL Nº: 456
    SECÇÃO: M
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 2,5
    TITULARES
    """

    pages = discover_caderneta_pages(early_contract_pages + caderneta)
    values = parse_caderneta(pages)

    assert [page.page_number for page in pages] == [19, 20]
    assert values.property_name == "Herdade da Fonte"
    assert values.matrix_article == "456"
    assert values.matrix_section == "M"
    assert values.area_m2 == 25000


def test_discovers_and_parses_caderneta_in_document_tail() -> None:
    front_pages = "\n".join(f"[Page {index}] contrato" for index in range(1, 11))
    caderneta = """
    [Page 11] CADERNETA PREDIAL RÚSTICA
    IDENTIFICAÇÃO DO PRÉDIO
    NOME/LOCALIZAÇÃO PRÉDIO: Quinta da Ribeira
    ARTIGO MATRICIAL Nº: 123
    SECÇÃO: K
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 1,481200
    TITULARES
    """

    pages = discover_caderneta_pages(front_pages + caderneta)
    values = parse_caderneta(pages)

    assert [page.page_number for page in pages] == [11]
    assert values.property_name == "Quinta da Ribeira"
    assert values.matrix_article == "123"
    assert values.matrix_section == "K"
    assert values.area_m2 == 14812


def test_reconciles_contract_and_caderneta_with_evidence() -> None:
    contract = PropertyExtractionPipeline().extract("""
    CONTRATO DE ARRENDAMENTO
    Senhorio, Arrendatário e renda mensal acordada.
    Considerando que:
    a) O prédio rústico denominado por Quinta da Ribeira, composto por olival,
    encontra-se inscrito na matriz sob o artigo 123, secção K.
    b) Cláusula seguinte.
    """)
    annex = """
    [Page 11] CADERNETA PREDIAL RÚSTICA
    IDENTIFICAÇÃO DO PRÉDIO
    NOME/LOCALIZAÇÃO PRÉDIO: Quinta da Ribeira
    ARTIGO MATRICIAL Nº: 123
    SECÇÃO: K
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 1,481200
    TITULARES
    """

    result = recover_from_annexes(contract, annex)

    assert result.area_m2 == 14812
    assert result.confidence == 100
    assert result.evidence_model["area_m2"]["source"] == "caderneta_predial"
    assert result.evidence_model["matrix_article"]["source"] == "contract_and_caderneta"
    assert result.audit["recovered_fields"] == ["area_m2"]


def test_property_output_uses_a_dedicated_worksheet_in_the_main_workbook() -> None:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise unittest.SkipTest("openpyxl is not installed in this environment")

    with TemporaryDirectory() as directory:
        workbook_path = Path(directory) / "document_register.xlsx"
        property_register = PropertyExcelRegister(workbook_path)
        property_register.append({
            "processed_at": "2026-08-28T00:00:00+00:00",
            "source_file_name": "lease.pdf",
            "sha256": "property-hash",
            "status": "processed",
            "property_name": "Quinta da Ribeira",
            "area_m2": 45230,
        })

        candidate = PdfCandidate(
            source_path=Path(directory) / "main.pdf",
            copied_path=Path(directory) / "main.pdf",
            sha256="main-hash",
            created_at=datetime.now(timezone.utc),
            modified_at=datetime.now(timezone.utc),
        )
        ExcelRegister(workbook_path).append(candidate, ExtractionResult())

        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path)
        assert PROPERTY_SHEET_NAME in workbook.sheetnames
        assert "Document Register" in workbook.sheetnames
        sheet = workbook[PROPERTY_SHEET_NAME]
        assert sheet.cell(row=2, column=3).value == "lease.pdf"
        assert property_register.existing_hashes() == {"property-hash"}
        workbook.close()
