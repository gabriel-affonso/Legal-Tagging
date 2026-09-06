from dataclasses import replace
import json
from pathlib import Path
from unittest.mock import patch

import pymupdf as fitz
import pytest

from doc_register.step50.extraction import deterministic_candidates, make_candidate, resolve, segment_pages
from doc_register.step50.schema import Region, Page, normalize
from doc_register.step50.pipeline import Pipeline, PipelineConfig
from doc_register.step50.evaluation import evaluate, validate_manifest
from doc_register.step50.export import export_results
from doc_register.step50.review import apply_decisions
from doc_register.step50.models import LocalModels


def region(text, method='native', logical='contract-p1', source='contract'):
    return Region('r1', 1, [0, 0, 300, 300], text, method, source, logical)


def pdf(path, pages):
    with fitz.open() as doc:
        for text in pages:
            page = doc.new_page()
            page.insert_text((40, 60), text, fontsize=11)
        doc.save(path)


def test_labelled_nif_stays_with_its_person():
    r = region('Senhorio: Ana Maria Lopes NIF 501964843\nArrendatário: Bruno Silva NIF 502011378')
    fields = resolve(deterministic_candidates(r))
    assert {(f.entity, f.value) for f in fields if f.field == 'tax_id'} == {('ana maria lopes', '501964843'), ('bruno silva', '502011378')}
    assert all(f.status == 'accepted' for f in fields)


def test_no_checksum_vote_for_wrong_entity():
    r = region('Senhorio: Ana Maria Lopes NIF 501964843\nArrendatário: Bruno Silva NIF 501964843')
    fields = resolve(deterministic_candidates(r))
    assert all(f.status == 'conflict' and f.value is None for f in fields if f.field == 'tax_id')


def test_roles_are_not_names_and_cadastral_is_not_lessor():
    assert not deterministic_candidates(region('Senhorio: Primeiro Outorgante'))[0].value
    assert not deterministic_candidates(region('Senhorio: Ana Maria Lopes', source='caderneta'))


def test_annual_per_hectare_preserved():
    candidates = deterministic_candidates(region('Renda anual de 1.200,50 EUR por hectare'))
    rent = next(c for c in candidates if c.field == 'rent')
    assert rent.value == {'amount': '1200.5', 'currency': 'EUR', 'frequency': 'annual', 'basis': 'per_ha', 'condition': ''}
    assert resolve(candidates)[0].status == 'needs_review'


def test_three_letter_section_not_truncated():
    candidates = deterministic_candidates(region('Artigo matricial: 154; Secção: ABC'))
    assert next(c.value for c in candidates if c.field == 'property_section') == 'ABC'


def test_notarial_signature_withheld():
    fields = resolve(deterministic_candidates(region('Reconhecimento notarial. Assinado em 20/05/2024')))
    assert fields[0].status == 'needs_review'


def test_conflicts_not_majority_voted_and_periods_separate():
    r = region('Assinado em 20/05/2024. Assinado em 21/05/2024.')
    candidates = deterministic_candidates(r)
    assert resolve(candidates + candidates)[0].status == 'conflict'
    other = region('Assinado em 22/05/2024', logical='amendment-p3', source='amendment')
    assert len(resolve(candidates + deterministic_candidates(other))) == 2


def test_visual_ambiguity_and_ocr_not_autoaccepted():
    r = region('Assinado em 20/05/2024', 'visual')
    r.issues = ['crossed_out', 'uncertain_tokens']
    assert resolve(deterministic_candidates(r))[0].status == 'needs_review'
    assert resolve(deterministic_candidates(region(r.text, 'ocr')))[0].status == 'needs_review'


def test_multiple_cadernetas_never_share_logical_identity():
    pages = [Page(i, 100, 100, 0, 'processed', [region('CADERNETA PREDIAL RUSTICA')]) for i in [1, 2]]
    assert len({s.logical_document for s in segment_pages(pages, 6000)}) == 2


def test_page_13_is_extracted_and_resume_is_idempotent(tmp_path):
    source = tmp_path / 'contract.pdf'
    pdf(source, ['CONTRATO DE ARRENDAMENTO'] + ['Clausula sem dados.'] * 11 + ['Assinado em 20/05/2024'])
    config = PipelineConfig(text_enabled=False, vision_enabled=False, ocr_enabled=False, native_min_chars=1)
    pipeline = Pipeline(tmp_path / 'output', config)
    path = pipeline.process(source)
    result = json.loads(path.read_text())
    assert len(result['pages']) == 13
    assert any(f['field'] == 'signed_date' and f['value'] == '2024-05-20' for f in result['fields'])
    assert pipeline.process(source) == path
    bounded = Pipeline(tmp_path / 'output', replace(config, max_pages=12)).process(source)
    limited = json.loads(bounded.read_text())
    assert limited['incomplete_coverage'] and limited['pages'][-1]['status'] == 'not_processed'
    assert not limited['fully_approved']


def test_hybrid_page_proposes_visual_even_with_native_text(tmp_path):
    source = tmp_path / 'hybrid.pdf'
    with fitz.open() as doc:
        p = doc.new_page()
        p.insert_text((20, 30), 'CONTRATO DE ARRENDAMENTO ' * 5)
        p.draw_line((20, 200), (120, 210))
        doc.save(source)
    config = PipelineConfig(text_enabled=False, vision_enabled=False, ocr_enabled=False)
    result = json.loads(Pipeline(tmp_path/'out', config).process(source).read_text())
    crops = [r for r in result['pages'][0]['regions'] if r['image']]
    assert crops and Path(crops[0]['image']['path']).exists()
    assert crops[0]['image']['width'] * crops[0]['image']['height'] <= config.max_pixels
    assert result['incomplete_coverage']


def test_partial_model_failure_keeps_native_evidence(tmp_path):
    source = tmp_path/'test.pdf'
    pdf(source, ['CONTRATO DE ARRENDAMENTO\nSenhorio: Ana Maria Lopes\nAssinado em 20/05/2024'])
    with patch.object(LocalModels, 'extract', side_effect=TimeoutError('test timeout')):
        result = json.loads(Pipeline(tmp_path/'out', PipelineConfig(vision_enabled=False, native_min_chars=1)).process(source).read_text())
    assert result['incomplete_coverage']
    assert any(f['status'] == 'accepted' for f in result['fields'])
    assert any('extraction_failure' in i for i in result['issues'])


def test_review_idempotency_preservation_and_export(tmp_path):
    source = tmp_path/'test.pdf'
    pdf(source, ['CONTRATO DE ARRENDAMENTO\nSenhorio: Ana Maria Lopes\nAssinado em 20/05/2024'])
    pipeline = Pipeline(tmp_path/'out', PipelineConfig(text_enabled=False, vision_enabled=False, native_min_chars=1))
    path = pipeline.process(source)
    result = json.loads(path.read_text())
    f = next(f for f in result['fields'] if f['field'] == 'signed_date')
    decisions = tmp_path/'decisions.json'
    decisions.write_text(json.dumps({'sha256': result['document']['sha256'], 'run_id': result['run_id'], 'decisions': [
        {**{k: f[k] for k in ['field', 'entity', 'logical_document']}, 'candidate_id': f['candidate_ids'][0], 'reason': 'Conferido no original'}]}))
    reviewed = apply_decisions(path, decisions, 'Test Reviewer')
    assert apply_decisions(path, decisions, 'Test Reviewer') == reviewed
    before = reviewed.read_bytes()
    assert pipeline.process(source, force=True) != path
    assert reviewed.read_bytes() == before
    assert next(f for f in json.loads(path.read_text())['fields'] if f['field'] == 'signed_date')['acceptance'] == 'automatic'
    exported = export_results([reviewed], tmp_path/'export.xlsx', reference='MG-100')
    from openpyxl import load_workbook
    with exported.open('rb') as handle:
        workbook = load_workbook(handle)
        assert {'Entrada_Rapida', 'db.Contrato', 'audit.Step50', 'Document Register'} <= set(workbook.sheetnames)
        workbook.close()
    with pytest.raises(FileExistsError):
        export_results([reviewed], exported)


def manifest_and_prediction():
    fact = {'field': 'signed_date', 'entity': 'contract', 'logical_document': 'contract-p1', 'value': '2024-05-20', 'legibility': 'legible', 'evidence_pages': [1]}
    manifest = {'documents': [{'id': 'doc', 'group_id': 'group', 'split': 'test', 'sha256': 'a'*64, 'annotators': ['human'], 'fully_annotated_fields': ['signed_date'], 'tags': ['digital'], 'facts': [fact]}]}
    predicted = {'document': {'sha256': 'a'*64}, 'fields': [{**fact, 'status': 'accepted', 'acceptance': 'automatic', 'candidate_ids': ['c']}],
                 'pages': [{'regions': [{'id': 'r', 'page': 1, 'text': 'Assinado em 20/05/2024'}]}],
                 'candidates': [{'id': 'c', 'region_id': 'r', 'value': '2024-05-20', 'quote': 'Assinado em 20/05/2024'}]}
    return manifest, predicted


def test_evaluator_counts_missing_duplicate_extra_and_evidence_errors():
    manifest, prediction = manifest_and_prediction()
    assert evaluate(manifest, {'doc': prediction})['micro']['recall'] == 1
    assert evaluate(manifest, {})['micro']['recall'] == 0
    prediction['fields'].append(dict(prediction['fields'][0]))
    assert evaluate(manifest, {'doc': prediction})['micro']['precision'] == .5
    prediction['candidates'][0]['quote'] = 'made up'
    report = evaluate(manifest, {'doc': prediction})
    assert report['micro']['recall'] == 0
    assert report['micro']['evidence_errors'] == 1


def test_evaluator_separates_human_and_automatic():
    manifest, prediction = manifest_and_prediction()
    prediction['fields'][0]['acceptance'] = 'human'
    assert evaluate(manifest, {'doc': prediction})['micro']['recall'] == 0
    assert evaluate(manifest, {'doc': prediction}, human=True)['micro']['recall'] == 1


def test_evaluator_rejects_leakage_and_hash_mismatch():
    manifest, prediction = manifest_and_prediction()
    manifest['documents'].append({**manifest['documents'][0], 'id': 'copy', 'split': 'development'})
    with pytest.raises(ValueError, match='leakage'):
        validate_manifest(manifest)
    manifest['documents'].pop()
    prediction['document']['sha256'] = 'b'*64
    with pytest.raises(ValueError, match='hash_mismatch'):
        evaluate(manifest, {'doc': prediction})


def test_strict_numbers_and_local_endpoint():
    with pytest.raises(ValueError):
        normalize('tax_id', '50196484X3')
    with pytest.raises(ValueError):
        normalize('signed_date', '31/02/2024')
    with pytest.raises(ValueError):
        LocalModels(PipelineConfig(ollama_url='https://example.com'))


def test_visual_runs_even_when_ocr_fails(tmp_path):
    from doc_register.step50.recognition import inventory_page
    with fitz.open() as document:
        document.new_page()
        with patch('doc_register.step50.recognition.ocr', side_effect=TimeoutError('ocr timeout')):
            page = inventory_page(document, 0, PipelineConfig(), tmp_path,
                                  lambda image: {'text': 'Assinado em 20/05/2024', 'modality': 'handwritten', 'uncertain_tokens': [], 'crossed_out': False})
    assert any(r.method == 'visual' and r.text for r in page.regions)
    assert any('ocr_failure' in issue for r in page.regions for issue in r.issues)


def test_visual_failure_is_visible_and_keeps_ocr(tmp_path):
    from doc_register.step50.recognition import inventory_page
    def unavailable(image):
        raise TimeoutError('visual timeout')
    with fitz.open() as document:
        document.new_page()
        with patch('doc_register.step50.recognition.ocr', return_value='Texto OCR válido'):
            page = inventory_page(document, 0, PipelineConfig(), tmp_path, unavailable)
    assert any(r.text == 'Texto OCR válido' for r in page.regions)
    assert any('visual_failure' in issue for r in page.regions for issue in r.issues)


def test_config_rejects_string_booleans():
    with pytest.raises(ValueError, match='bool'):
        PipelineConfig(ocr_enabled='false')
