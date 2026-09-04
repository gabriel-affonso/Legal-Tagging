from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
import sys
import unittest

from openpyxl import load_workbook

from doc_register.models import ExtractionResult, PdfCandidate
from doc_register.step40 import (
    ENTRY_COLUMNS,
    PROCESSING_DONE,
    PROCESSING_REVIEW,
    TABLE_SCHEMAS,
    NormalizedWorkbookRegister,
    valid_iban,
)


def _candidate(digest: str = "a" * 64) -> PdfCandidate:
    now = datetime.now(timezone.utc)
    return PdfCandidate(Path("MG-100_contract.pdf"), Path("/tmp/MG-100_contract.pdf"), digest, now, now)


def _result() -> ExtractionResult:
    return ExtractionResult(
        document_category="lease_contract",
        document_type="lease_contract",
        contract_type="ARRENDAMENTO",
        signed_date="2026-01-10",
        bank_account_holder="Ana Silva",
        iban="PT50000201231234567890154",
        raw_json={
            "step3_3_field_centric": {
                "cadastral_properties": [
                    {
                        "codigo_interno": "PR-10", "property_article": "42", "property_section": "A",
                        "property_parish": "Ala", "property_municipality": "Mogadouro",
                        "property_district": "Bragança", "property_total_area": "15000 m2",
                        "owner_name": "Ana Silva", "owner_tax_id": "501964843", "pages": [2],
                    },
                    {
                        "codigo_interno": "PR-11", "property_article": "43", "property_section": "B",
                        "property_parish": "Ala", "property_municipality": "Mogadouro",
                        "property_district": "Bragança", "property_total_area": "2 ha",
                        "owner_name": "Bruno Silva", "owner_tax_id": "502011378", "pages": [3],
                    },
                ]
            }
        },
    )


def test_step40_schema_relations_area_and_idempotency() -> None:
    assert valid_iban("PT50000201231234567890154")
    with TemporaryDirectory() as directory:
        path = Path(directory) / "normalized.xlsx"
        register = NormalizedWorkbookRegister(path)
        first = register.publish(_candidate(), _result(), reference_context="MG-100", batch_id="LOTE-A")
        second = register.publish(_candidate(), _result(), reference_context="MG-100", batch_id="LOTE-A")
        assert first[PROCESSING_DONE] == 2
        assert second[PROCESSING_DONE] == 0

        workbook = load_workbook(path, data_only=False)
        expected = ["INSTRUCOES", "Entrada_Rapida", *TABLE_SCHEMAS, "Listas"]
        assert workbook.sheetnames[:len(expected)] == expected
        assert [cell.value for cell in workbook["Entrada_Rapida"][1]] == list(ENTRY_COLUMNS)
        for sheet_name, columns in TABLE_SCHEMAS.items():
            sheet = workbook[sheet_name]
            assert [cell.value for cell in sheet[1]] == list(columns)
            assert sheet.tables
            assert next(iter(sheet.tables.values())).ref.endswith(str(sheet.max_row))

        assert workbook["db.Contrato"].max_row == 2
        assert workbook["db.Terreno"].max_row == 3
        assert workbook["db.Proprietario"].max_row == 3
        assert workbook["db.ContratoTerreno"].max_row == 3
        assert workbook["db.ProprietarioTerreno"].max_row == 3
        assert workbook["db.ContaBancaria"].max_row == 2
        terrain_headers = {cell.value: index for index, cell in enumerate(workbook["db.Terreno"][1], start=1)}
        values = list(workbook["db.Terreno"].iter_rows(min_row=2, values_only=True))
        assert {row[terrain_headers["area_at_ha"] - 1] for row in values} == {1.5, 2}
        owner_terrain_headers = {cell.value: index for index, cell in enumerate(workbook["db.ProprietarioTerreno"][1], start=1)}
        parcel_ids = {row[terrain_headers["parcela_id"] - 1] for row in values}
        links = list(workbook["db.ProprietarioTerreno"].iter_rows(min_row=2, values_only=True))
        assert {row[owner_terrain_headers["terreno_id"] - 1] for row in links} == parcel_ids
        workbook.close()

        # A corrected/reopened staging row must reuse its natural IDs and must
        # not create duplicate entity or relation rows.
        workbook = load_workbook(path)
        entry_headers = {cell.value: index for index, cell in enumerate(workbook["Entrada_Rapida"][1], start=1)}
        original_ids = [
            workbook["Entrada_Rapida"].cell(2, entry_headers[column]).value
            for column in ("contrato_id", "proprietario_id", "terreno_id")
        ]
        workbook["Entrada_Rapida"].cell(2, entry_headers["estado_processamento"]).value = "PENDENTE"
        workbook.save(path)
        workbook.close()
        assert register.process_pending()[PROCESSING_DONE] == 1
        workbook = load_workbook(path)
        assert [workbook["Entrada_Rapida"].cell(2, entry_headers[column]).value for column in ("contrato_id", "proprietario_id", "terreno_id")] == original_ids
        assert workbook["db.Contrato"].max_row == 2
        assert workbook["db.Terreno"].max_row == 3
        assert workbook["db.ProprietarioTerreno"].max_row == 3
        workbook.close()


def test_step40_missing_reference_is_auditable_but_not_published() -> None:
    with TemporaryDirectory() as directory:
        path = Path(directory) / "normalized.xlsx"
        result = _result()
        register = NormalizedWorkbookRegister(path)
        register.publish(PdfCandidate(Path("contract.pdf"), Path("/tmp/contract.pdf"), "b" * 64, datetime.now(timezone.utc), datetime.now(timezone.utc)), result)
        workbook = load_workbook(path, data_only=False)
        entry_headers = {cell.value: index for index, cell in enumerate(workbook["Entrada_Rapida"][1], start=1)}
        assert workbook["Entrada_Rapida"].cell(2, entry_headers["estado_processamento"]).value == PROCESSING_REVIEW
        assert workbook["db.Contrato"].max_row == 1
        conflicts = list(workbook["audit.Conflicts"].iter_rows(min_row=2, values_only=True))
        assert any("referencia_mg ausente" in row[5] for row in conflicts)
        workbook.close()


def load_tests(loader, tests, pattern):
    module = sys.modules[__name__]
    functions = [
        getattr(module, name)
        for name in sorted(dir(module))
        if name.startswith("test_") and callable(getattr(module, name))
    ]
    return unittest.TestSuite(unittest.FunctionTestCase(function) for function in functions)
