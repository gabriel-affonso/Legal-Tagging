from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import sys
import json
from tempfile import TemporaryDirectory
import unittest

from doc_register.__main__ import _build_parser
from doc_register.models import ExtractionResult, PdfCandidate
from doc_register.property_table import build_property_table_rows
from doc_register.registry import (
    PROPERTY_TABLE_COLUMNS,
    PROPERTY_REGISTER_COLUMNS,
    PROPERTY_SHEET_NAME,
    PROPERTY_TABLE_SHEET_NAME,
    PropertyTableRegister,
)
from doc_register.step33_engine import prepare_field_centric_extraction


def _result(properties: list[dict[str, object]]) -> ExtractionResult:
    return ExtractionResult(
        document_category="lease_contract",
        document_type="Contrato de Arrendamento",
        lessor="Ana Maria Lopes; Augusto Mendes",
        lessee="GESTO ENERGIA, S.A.",
        lessee_tax_id="508567475",
        rent_amount="1000.00",
        rent_currency="EUR",
        rent_frequency="annual",
        rent_unit="per_hectare",
        raw_json={
            "step3_3_field_centric": {
                "cadastral_properties": properties,
            }
        },
    )


def _properties() -> list[dict[str, object]]:
    return [
        {
            "matrix_key": "4-L",
            "property_name": "Serra da Abelha",
            "property_article": "4",
            "property_section": "L",
            "property_parish": "Penas Roias",
            "property_municipality": "Mogadouro",
            "property_district": "Bragança",
            "property_total_area": "76.775 hectares",
            "property_total_area_m2": 767750,
            "owner_name": "FREGUESIA DE PENAS ROIAS",
            "owner_tax_id": "508574935",
            "pages": [17, 18],
            "matched_contract_identity": True,
        },
        {
            "matrix_key": "29-F",
            "property_name": "Serra da Senhora",
            "property_article": "29",
            "property_section": "F",
            "property_parish": "Penas Roias",
            "property_municipality": "Mogadouro",
            "property_district": "Bragança",
            "property_total_area": "48.8811 hectares",
            "property_total_area_m2": 488811,
            "owner_name": "FREGUESIA DE PENAS ROIAS",
            "owner_tax_id": "508574935",
            "pages": [19, 20],
            "matched_contract_identity": True,
        },
    ]


def test_multi_property_rows_repeat_contract_and_keep_property_identity_atomic() -> None:
    rows = build_property_table_rows(_result(_properties()))

    assert len(rows) == 2
    assert [row["property_matrix_key"] for row in rows] == ["4-L", "29-F"]
    assert [row["property_name"] for row in rows] == ["Serra da Abelha", "Serra da Senhora"]
    assert [row["property_total_area_m2"] for row in rows] == [767750, 488811]
    assert all(row["lessee"] == "GESTO ENERGIA, S.A." for row in rows)
    assert all(row["lessor"] == "Ana Maria Lopes; Augusto Mendes" for row in rows)
    assert all(row["property_count"] == 2 for row in rows)
    assert all(row["leased_parcel_area"] == "" for row in rows)


def test_contract_matrix_matches_exclude_unrelated_internal_cadernetas() -> None:
    properties = _properties()
    properties[1] = {**properties[1], "matched_contract_identity": False}

    rows = build_property_table_rows(_result(properties))

    assert len(rows) == 1
    assert rows[0]["property_matrix_key"] == "4-L"
    assert rows[0]["property_match_status"] == "contract_identity_match"


def test_unconfirmed_internal_properties_are_kept_but_require_review() -> None:
    properties = [
        {**item, "matched_contract_identity": False}
        for item in _properties()
    ]

    rows = build_property_table_rows(_result(properties))

    assert len(rows) == 2
    assert all(row["property_match_status"] == "internal_annex_unconfirmed" for row in rows)
    assert all(row["needs_review"] == "yes" for row in rows)
    assert all("property_contract_identity_unconfirmed" in row["review_reason"] for row in rows)


def test_missing_property_creates_one_explicit_unresolved_audit_row() -> None:
    rows = build_property_table_rows(_result([]))

    assert len(rows) == 1
    assert rows[0]["property_count"] == 0
    assert rows[0]["property_row_status"] == "unresolved_no_property"
    assert rows[0]["property_matrix_key"] == ""
    assert rows[0]["needs_review"] == "yes"


def test_non_contract_attachments_do_not_create_property_table_rows() -> None:
    bank_attachment = ExtractionResult(
        document_category="bank_details",
        property_article="140",
        property_section="J",
    )

    assert build_property_table_rows(bank_attachment) == []


def test_property_extraction_is_reused_before_lower_priority_main_sources() -> None:
    result = _result([
        {**_properties()[0], "matrix_key": "999-Z", "property_article": "999", "property_section": "Z"}
    ])
    result.raw_json["property_extraction"] = {
        "properties": [
            {**_properties()[0], "matrix_key": "33-B", "property_article": "33", "property_section": "B", "matrix_article": "33", "matrix_section": "B"},
            {**_properties()[1], "matrix_key": "61-B", "property_article": "61", "property_section": "B", "matrix_article": "61", "matrix_section": "B"},
        ]
    }

    rows = build_property_table_rows(result)

    assert [row["property_matrix_key"] for row in rows] == ["33-B", "61-B"]


def test_duplicate_matrix_groups_merge_pages_and_keep_the_complete_values() -> None:
    duplicate = {
        "matrix_key": "4-L",
        "property_article": "4",
        "property_section": "L",
        "pages": [21],
        "matched_contract_identity": True,
    }
    rows = build_property_table_rows(_result([_properties()[0], duplicate]))

    assert len(rows) == 1
    assert rows[0]["property_name"] == "Serra da Abelha"
    assert rows[0]["property_source_pages"] == "17; 18; 21"


def test_invalid_generic_property_name_is_not_published_as_a_fact() -> None:
    properties = [{
        **_properties()[0],
        "property_name": "Encargos potenciais sobre o prédio",
    }]

    row = build_property_table_rows(_result(properties))[0]

    assert row["property_name"] == ""
    assert row["needs_review"] == "yes"
    assert "property_name_invalid_or_missing" in row["review_reason"]


def test_property_table_cli_flag_is_available_for_main_scan() -> None:
    args = _build_parser().parse_args(["scan", "--property-table"])
    assert args.command == "scan"
    assert args.property_table is True


def test_step33_structured_properties_include_locality_and_numeric_area() -> None:
    text = """[Page 1]
    CONTRATO DE ARRENDAMENTO
    Considerando que: artigo 4, secção L; artigo 29, secção F.
    Cláusula 1
    [Page 17]
    CADERNETA PREDIAL RÚSTICA
    Modelo A
    IDENTIFICAÇÃO DO PRÉDIO
    ARTIGO MATRICIAL: 4
    SECÇÃO: L
    NOME/LOCALIZAÇÃO PRÉDIO: Serra da Abelha
    Freguesia: Penas Roias
    Concelho: Mogadouro
    Distrito: Bragança
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 76,775000
    [Page 19]
    CADERNETA PREDIAL RÚSTICA
    Modelo B
    IDENTIFICAÇÃO DO PRÉDIO
    ARTIGO MATRICIAL: 29
    SECÇÃO: F
    NOME/LOCALIZAÇÃO PRÉDIO: Serra da Senhora
    Freguesia: Penas Roias
    Concelho: Mogadouro
    Distrito: Bragança
    ELEMENTOS DO PRÉDIO
    ÁREA TOTAL (HA): 48,881100
    """
    properties = prepare_field_centric_extraction(text).cadastral_properties

    assert [item["matrix_key"] for item in properties] == ["4-L", "29-F"]
    assert all(item["property_municipality"] == "Mogadouro" for item in properties)
    assert all(item["property_district"] == "Bragança" for item in properties)
    assert [item["property_total_area_m2"] for item in properties] == [767750, 488811]


def test_property_table_upsert_replaces_all_rows_for_contract() -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return

    with TemporaryDirectory() as directory:
        workbook_path = Path(directory) / "register.xlsx"
        now = datetime.now(timezone.utc)
        candidate = PdfCandidate(
            source_path=Path("PR100_CA.pdf"),
            copied_path=Path(directory) / "processing" / "PR100_CA.pdf",
            sha256="a" * 64,
            created_at=now,
            modified_at=now,
        )
        register = PropertyTableRegister(workbook_path)
        assert register.upsert(candidate, _result(_properties())) == 2
        assert register.upsert(candidate, _result([_properties()[1]])) == 1

        workbook = load_workbook(workbook_path)
        sheet = workbook[PROPERTY_TABLE_SHEET_NAME]
        headers = [cell.value for cell in sheet[1]]
        assert headers == list(PROPERTY_TABLE_COLUMNS)
        assert sheet.max_row == 2
        assert sheet.cell(2, headers.index("property_matrix_key") + 1).value == "29-F"
        assert sheet.cell(2, headers.index("property_row_id") + 1).value == f"{'a' * 64}::29-F"
        assert sheet.cell(2, headers.index("property_total_area_m2") + 1).value == 488811
        assert sheet.tables["PropertyTable"].ref.endswith("2")
        workbook.close()


def test_register_enriches_main_result_with_independent_property_extraction() -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return

    with TemporaryDirectory() as directory:
        workbook_path = Path(directory) / "register.xlsx"
        digest = "b" * 64
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = PROPERTY_SHEET_NAME
        sheet.append(PROPERTY_REGISTER_COLUMNS)
        payload = {column: "" for column in PROPERTY_REGISTER_COLUMNS}
        payload.update({
            "sha256": digest,
            "document_type": "lease_contract",
            "status": "processed_multi_property",
            "properties": json.dumps([
                {**_properties()[0], "matrix_key": "33-B", "property_article": "33", "property_section": "B", "matrix_article": "33", "matrix_section": "B"},
                {**_properties()[1], "matrix_key": "61-B", "property_article": "61", "property_section": "B", "matrix_article": "61", "matrix_section": "B"},
            ]),
        })
        sheet.append([payload[column] for column in PROPERTY_REGISTER_COLUMNS])
        workbook.save(workbook_path)
        workbook.close()
        now = datetime.now(timezone.utc)
        candidate = PdfCandidate(Path("PR100_CA.pdf"), Path("PR100_CA.pdf"), digest, now, now)
        main_result = ExtractionResult(document_category="other")

        enriched = PropertyTableRegister(workbook_path).enrich_from_property_extraction(candidate, main_result)
        rows = build_property_table_rows(enriched)

    assert enriched.document_category == "lease_contract"
    assert [row["property_matrix_key"] for row in rows] == ["33-B", "61-B"]
    assert all(row["property_row_status"] == "resolved" for row in rows)


def load_tests(loader, tests, pattern):
    module = sys.modules[__name__]
    functions = [
        getattr(module, name)
        for name in sorted(dir(module))
        if name.startswith("test_") and callable(getattr(module, name))
    ]
    return unittest.TestSuite(unittest.FunctionTestCase(function) for function in functions)
