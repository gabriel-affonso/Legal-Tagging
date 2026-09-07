"""Executable inventory of every existing schema column; unknown semantics stay explicit."""
from ..schemas import REGISTER_COLUMNS,TECHNICAL_FIELDS,OPERATIONAL_FIELDS
from ..step40 import ENTRY_COLUMNS,TABLE_SCHEMAS
from ..registry import PROPERTY_REGISTER_COLUMNS, PROPERTY_TABLE_COLUMNS

MAP={
    'name':('name','person|organization'), 'nome_canonico':('name','person|organization'),
    'nome_proprietario':('name','person|organization'), 'nif':('tax_id','person|organization'),
    'papel_no_contrato':('contract_role','participation'), 'papel':('contract_role','participation'),
    'artigo_matricial':('cadastral_article','property'), 'property_article':('cadastral_article','property'),
    'descricao_predial':('registry_description','property'), 'registry_description':('registry_description','property'),
    'freguesia':('property_parish','property'), 'concelho':('property_municipality','property'),
    'distrito':('property_district','property'), 'local':('property_name','property'),
    'data_assinatura':('signed_date','contract'), 'signed_date':('signed_date','contract'),
    'area_at_m2':('area_measurement[cadastral_total,m2]','property'), 'area_at_ha':('convert_area_to_ha[cadastral_total]','property'),
    'area_medida_m2':('area_measurement[measured,m2]','property'), 'area_medida_ha':('convert_area_to_ha[measured]','property'),
    'area_documental_parcela_ha':('area_measurement[leased_usable,ha,explicit_property_scope]','property|lease_object'),
    'area_documental_total_ha':('area_measurement[leased_usable,ha]','lease_object'),
    'area_considerada_ha':('area_measurement[considered,ha]','property|lease_object'),
    'iban':('iban','person|organization'), 'payment_date':('bank_execution_date','payment'),
    'payment_reference':('payment_reference','payment'),
    'payment_gross_amount':('payment_gross_amount','payment'), 'withholding_tax_amount':('withholding_tax_amount','payment'),
    'payment_net_amount':('payment_net_amount','payment'), 'contract_end_date':('contract_end_date','contract'),
}
MAP.update({name:(name,'property') for name in ('property_name','property_parish','property_municipality','property_district')})
MAP.update({'matrix_article':('cadastral_article','property'),'matrix_section':('cadastral_section','property'),
    'property_section':('cadastral_section','property'),'owner_name':('ownership.holder.name','person|organization'),
    'owner_tax_id':('ownership.holder.tax_id','person|organization'),
    'payer':('relations[payer]','payment'),'payee':('relations[payee]','payment'),
    'lessor':('relations[has_role,lessor]','contract'),'lessor_2':('relations[has_role,lessor]','contract'),
    'lessee':('relations[has_role,lessee]','contract'),
    'leased_parcel_area':('area_measurement[leased_usable,explicit_scope]','property|lease_object')})
OPERATIONS={'referencia_mg','referencia_mg*','codigo_interno','projeto','lote','entrada_id','contrato_id','proprietario_id','terreno_id','parcela_id'}


def field_catalog(workbook=None):
    tables={'Document Register':REGISTER_COLUMNS,'Property Extraction':PROPERTY_REGISTER_COLUMNS,
            'Property Table':PROPERTY_TABLE_COLUMNS,'Entrada_Rapida':ENTRY_COLUMNS,**TABLE_SCHEMAS}
    source='repository_schemas_not_live_workbook'
    workbook_hash=None
    if workbook:
        from openpyxl import load_workbook
        from .storage import digest_file
        from pathlib import Path
        book=load_workbook(workbook,read_only=True,data_only=False)
        try:tables={s.title:[str(x) for x in next(s.iter_rows(min_row=1,max_row=1,values_only=True)) if x is not None] for s in book}
        finally:book.close()
        source=str(Path(workbook).resolve());workbook_hash=digest_file(Path(workbook))
    rows=[]
    for table,columns in tables.items():
        for column in columns:
            mapping=MAP.get(column)
            technical=column in set(TECHNICAL_FIELDS)|set(OPERATIONAL_FIELDS)|OPERATIONS
            origins=['operacional'] if technical else ['documental','derivado','externo','operacional','humano']
            rows.append({'table':table,'column':column,'canonical_path':mapping[0] if mapping else None,
                'entity_type':mapping[1] if mapping else None,'status':'mapped' if mapping else 'operational' if technical else 'unmapped',
                'allowed_origins':origins,'overwrite_policy':'never_overwrite_existing_operational_or_human_values',
                'value_origin_policy':'per_assertion_provenance_required_not_inferred_from_column',
                'scope_requirement':'explicit_property_scope_no_joint_area_copy' if column=='area_documental_parcela_ha' else 'entity_and_period_resolved',
                'export_eligibility':'accepted_and_scope_resolved' if mapping else 'not_projectable',
                'note':'Scope and source are mandatory; conversion keeps input assertion IDs.'})
    return {'schema_version':'5.5.1','inventory_source':source,'workbook_sha256':workbook_hash,'columns':rows}
