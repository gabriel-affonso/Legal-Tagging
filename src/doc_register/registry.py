from __future__ import annotations

from datetime import datetime, timezone
import json
from pathlib import Path
from contextlib import contextmanager

try:
    import fcntl
except ImportError:  # pragma: no cover - Windows fallback
    fcntl = None

from .models import ExtractionResult, PdfCandidate, REGISTER_COLUMNS
from .property_table import build_property_table_rows
from .schemas import OPERATIONAL_FIELDS


SHEET_NAME = "Document Register"
PROPERTY_SHEET_NAME = "Property Extraction"
PROPERTY_TABLE_SHEET_NAME = "Property Table"

PROPERTY_TABLE_METADATA_COLUMNS = (
    "property_table_version",
    "property_row_id",
    "property_index",
    "property_count",
    "property_row_status",
    "property_match_status",
    "property_association_status",
    "property_association_evidence",
    "contract_operational_property_id",
    "property_source_pages",
    "property_structure_status",
    "property_owner_status",
    "property_source_file",
    "property_total_area_m2",
    "property_json",
)
PROPERTY_TABLE_COLUMNS = tuple([
    *OPERATIONAL_FIELDS,
    *PROPERTY_TABLE_METADATA_COLUMNS,
    *(column for column in REGISTER_COLUMNS if column not in OPERATIONAL_FIELDS),
])

PROPERTY_REGISTER_COLUMNS = (
    "pipeline_version",
    "processed_at",
    "source_file_name",
    "source_file_path",
    "sha256",
    "status",
    "reason",
    "document_type",
    "owner_name",
    "property_name",
    "matrix_article",
    "matrix_section",
    "property_matrix_key",
    "area_m2",
    "confidence",
    "extraction_confidence",
    "identity_confidence",
    "completeness_score",
    "consistency_score",
    "final_confidence",
    "properties",
    "lease_score",
    "candidate_score",
    "used_llm",
    "source_section",
    "source_clause",
    "text_source",
    "native_text_chars",
    "ocr_text_chars",
    "extraction_notes",
    "annex_text_chars",
    "llm_error",
    "evidence_model",
    "property_catalog_status",
    "property_catalog_text_source",
    "property_pack_status",
    "property_pack_match_score",
    "property_pack_match_method",
    "property_pack_candidate_count",
    "property_pack_property_count",
    "caderneta_source_file",
    "crp_source_file",
    "caderneta_evidence_sources",
    "caderneta_same_pdf_found",
    "internal_caderneta_status",
    "internal_caderneta_count",
    "audit",
)


class ExcelRegister:
    def __init__(self, path: Path):
        self.path = path

    def existing_hashes(self) -> set[str]:
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            sha_col = REGISTER_COLUMNS.index("sha256") + 1
            values: set[str] = set()
            for row in range(2, sheet.max_row + 1):
                value = sheet.cell(row=row, column=sha_col).value
                if value:
                    values.add(str(value))
            workbook.close()
            return values

    def append(self, candidate: PdfCandidate, result: ExtractionResult) -> None:
        self._write(candidate, result, replace_existing=False)

    def recent_candidates(self, limit: int) -> list[PdfCandidate]:
        """Return recoverable records newest-first for Step 3.6 batch recall."""
        if limit <= 0:
            return []
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            headers = {str(sheet.cell(row=1, column=index).value): index for index in range(1, sheet.max_column + 1)}
            required = {"source_file_name", "copied_file_path", "sha256"}
            if not required.issubset(headers):
                workbook.close()
                return []
            candidates: list[PdfCandidate] = []
            for row in range(sheet.max_row, 1, -1):
                copied = Path(str(sheet.cell(row=row, column=headers["copied_file_path"]).value or ""))
                source_name = str(sheet.cell(row=row, column=headers["source_file_name"]).value or copied.name)
                if not copied.is_file():
                    continue
                digest = str(sheet.cell(row=row, column=headers["sha256"]).value or "").strip()
                if not digest:
                    continue
                modified = datetime.fromtimestamp(copied.stat().st_mtime, tz=timezone.utc)
                # The original may no longer be present in the intake folder;
                # retain its filename for audit while processing the immutable
                # copied PDF.
                candidates.append(PdfCandidate(Path(source_name), copied, digest, modified, modified))
                if len(candidates) >= limit:
                    break
            workbook.close()
            return candidates

    def upsert(self, candidate: PdfCandidate, result: ExtractionResult) -> None:
        """Replace the registered row for a PDF hash, or add it when absent."""
        self._write(candidate, result, replace_existing=True)

    def _write(
        self,
        candidate: PdfCandidate,
        result: ExtractionResult,
        *,
        replace_existing: bool,
    ) -> None:
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            operational = {
                "processed_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
                "source_file_name": candidate.source_path.name,
                "copied_file_path": str(candidate.copied_path),
                "sha256": candidate.sha256,
                "file_created_at": candidate.created_at.isoformat(),
                "file_modified_at": candidate.modified_at.isoformat(),
            }
            row = []
            for column in REGISTER_COLUMNS:
                value = operational.get(column, getattr(result, column, ""))
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                row.append(value)
            sha_col = REGISTER_COLUMNS.index("sha256") + 1
            existing_row = next(
                (
                    row_index
                    for row_index in range(2, sheet.max_row + 1)
                    if str(sheet.cell(row=row_index, column=sha_col).value or "") == candidate.sha256
                ),
                None,
            ) if replace_existing else None
            if existing_row:
                for column_index, value in enumerate(row, start=1):
                    sheet.cell(row=existing_row, column=column_index, value=value)
            else:
                sheet.append(row)
            self._format(sheet)
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
            workbook.close()

    def _load(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError as exc:
            raise RuntimeError("Missing dependency: install openpyxl with `pip install -r requirements.txt`.") from exc

        if self.path.exists():
            workbook = load_workbook(self.path)
            sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.create_sheet(SHEET_NAME, 0)
            migrated = _ensure_headers(sheet)
            if not sheet.tables:
                _add_table(sheet, "DocumentRegister", len(REGISTER_COLUMNS))
                migrated = True
            if migrated:
                self._format(sheet)
                workbook.save(self.path)
            return workbook, sheet

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(REGISTER_COLUMNS)
        _add_table(sheet, "DocumentRegister", len(REGISTER_COLUMNS))
        self._format(sheet)
        return workbook, sheet

    def _format(self, sheet) -> None:
        sheet.freeze_panes = "A2"
        widths_by_header = {
            "processed_at": 22,
            "source_file_name": 34,
            "copied_file_path": 48,
            "sha256": 66,
            "processing_status": 20,
            "review_reason": 48,
            "error_message": 54,
            "document_category": 24,
            "document_type": 26,
            "document_date": 18,
            "summary": 54,
            "extraction_notes": 42,
            "ocr_quality_flags": 34,
            "lessor": 34,
            "lessor_2": 34,
            "lessee": 34,
            "lessee_tax_id": 18,
            "option_price": 18,
            "annual_rent": 18,
            "rent_amount": 18,
            "rent_frequency": 16,
            "rent_unit": 18,
            "purchase_price": 18,
            "assignment_price": 18,
            "property_display_name": 42,
            "property_address": 48,
            "owner_address": 48,
            "bank_account_holder": 34,
            "iban": 34,
            "payment_description": 54,
            "ai_review_status": 22,
            "ai_reviewed_fields": 34,
            "ai_accepted_fields": 34,
            "ai_rejected_fields": 34,
            "ai_review_reason": 54,
            "cadastral_evidence_status": 32,
            "cadastral_evidence_pages": 22,
            "cadastral_conflicts": 48,
            "deterministic_json": 80,
            "llm_json": 80,
            "raw_json": 80,
        }
        for index, header in enumerate(REGISTER_COLUMNS, start=1):
            width = widths_by_header.get(header)
            if width:
                sheet.column_dimensions[_column_letter(index)].width = width

        if sheet.tables:
            table = next(iter(sheet.tables.values()))
            _set_table_ref(sheet, table, len(REGISTER_COLUMNS))


class PropertyTableRegister:
    """Step 3.4 worksheet with one row per resolved cadastral property."""

    def __init__(self, path: Path):
        self.path = path

    def existing_hashes(self) -> set[str]:
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            sha_col = PROPERTY_TABLE_COLUMNS.index("sha256") + 1
            values = {
                str(sheet.cell(row=row, column=sha_col).value)
                for row in range(2, sheet.max_row + 1)
                if sheet.cell(row=row, column=sha_col).value
            }
            workbook.close()
            return values

    def eligible_contract_hashes(self) -> set[str]:
        """Return only PDFs already classified as lease contracts.

        ``property-scan`` is the eligibility gate for Property Table.  This
        prevents the heavier main pipeline from opening non-contract PDFs just
        to eventually discard them from the worksheet.
        """
        if not self.path.exists():
            return set()
        with _workbook_lock(self.path):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("Missing dependency: install openpyxl with `pip install -r requirements.txt`.") from exc
            workbook = load_workbook(self.path, read_only=True, data_only=True)
            try:
                if PROPERTY_SHEET_NAME not in workbook.sheetnames:
                    return set()
                sheet = workbook[PROPERTY_SHEET_NAME]
                headers = {
                    str(cell.value): index
                    for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=False)), start=1)
                    if cell.value
                }
                if not {"sha256", "document_type"}.issubset(headers):
                    return set()
                return {
                    str(values[headers["sha256"] - 1])
                    for values in sheet.iter_rows(min_row=2, values_only=True)
                    if str(values[headers["document_type"] - 1] or "").strip().lower() == "lease_contract"
                    and values[headers["sha256"] - 1]
                }
            finally:
                workbook.close()

    def materialize_from_property_extraction(self) -> tuple[int, int]:
        """Rebuild Property Table from workbook data without opening any PDF.

        ``Property Extraction`` supplies the eligible lease contracts and
        their property-grain facts. ``Document Register`` contributes the
        already-resolved contract fields. The complete replacement is done in
        one locked workbook write, so a table run never invokes OCR or Ollama.
        """
        if not self.path.exists():
            return 0, 0
        with _workbook_lock(self.path):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("Missing dependency: install openpyxl with `pip install -r requirements.txt`.") from exc
            workbook = load_workbook(self.path)
            try:
                if PROPERTY_SHEET_NAME not in workbook.sheetnames:
                    return 0, 0
                extraction_sheet = workbook[PROPERTY_SHEET_NAME]
                extraction_headers = _sheet_headers(extraction_sheet)
                required = {"sha256", "document_type", "properties"}
                if not required.issubset(extraction_headers):
                    return 0, 0
                register_rows = (
                    _sheet_rows_by_sha(workbook[SHEET_NAME])
                    if SHEET_NAME in workbook.sheetnames else {}
                )
                if PROPERTY_TABLE_SHEET_NAME in workbook.sheetnames:
                    table_sheet = workbook[PROPERTY_TABLE_SHEET_NAME]
                else:
                    table_sheet = workbook.create_sheet(PROPERTY_TABLE_SHEET_NAME)
                    table_sheet.append(PROPERTY_TABLE_COLUMNS)
                    _add_table(table_sheet, "PropertyTable", len(PROPERTY_TABLE_COLUMNS))
                migrated = _ensure_named_headers(table_sheet, PROPERTY_TABLE_COLUMNS)
                if not table_sheet.tables:
                    _add_table(table_sheet, "PropertyTable", len(PROPERTY_TABLE_COLUMNS))
                    migrated = True
                if table_sheet.max_row > 1:
                    table_sheet.delete_rows(2, table_sheet.max_row - 1)

                contract_count = 0
                row_count = 0
                for extraction in _sheet_row_dicts(extraction_sheet, extraction_headers):
                    sha256 = str(extraction.get("sha256") or "")
                    if not sha256 or str(extraction.get("document_type") or "").strip().lower() != "lease_contract":
                        continue
                    contract_count += 1
                    register_row = register_rows.get(sha256, {})
                    result = _materialized_result(register_row, extraction)
                    operational = _materialized_operational(register_row, extraction, sha256)
                    for property_row in build_property_table_rows(result):
                        payload = {**property_row, **operational}
                        payload["property_row_id"] = f"{sha256}::{payload.get('property_row_id') or 'unresolved'}"
                        table_sheet.append([
                            _excel_value(payload.get(column, ""))
                            for column in PROPERTY_TABLE_COLUMNS
                        ])
                        row_count += 1
                self._format(table_sheet)
                workbook.save(self.path)
                return contract_count, row_count
            finally:
                workbook.close()

    def enrich_from_property_extraction(
        self, candidate: PdfCandidate, result: ExtractionResult
    ) -> ExtractionResult:
        """Attach the independent Property Extraction result for this PDF.

        The secondary pipeline is the richer property source for many legacy
        contracts.  Reading it by immutable file hash keeps the main scan and
        the materialized table aligned without using filenames as identities.
        """
        if not self.path.exists():
            return result
        with _workbook_lock(self.path):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("Missing dependency: install openpyxl with `pip install -r requirements.txt`.") from exc
            workbook = load_workbook(self.path, read_only=True, data_only=True)
            try:
                if PROPERTY_SHEET_NAME not in workbook.sheetnames:
                    return result
                sheet = workbook[PROPERTY_SHEET_NAME]
                headers = {
                    str(cell.value): index
                    for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=False)), start=1)
                    if cell.value
                }
                required = {"sha256", "document_type", "status", "properties"}
                if not required.issubset(headers):
                    return result
                for values in sheet.iter_rows(min_row=2, values_only=True):
                    if str(values[headers["sha256"] - 1] or "") != candidate.sha256:
                        continue
                    document_type = str(values[headers["document_type"] - 1] or "").strip().lower()
                    status = str(values[headers["status"] - 1] or "").strip()
                    if document_type != "lease_contract" or status not in {"processed", "processed_multi_property", "needs_review"}:
                        return result
                    properties = _json_list(values[headers["properties"] - 1])
                    if not properties:
                        properties = [_property_extraction_single(values, headers)]
                    raw = dict(result.raw_json) if isinstance(result.raw_json, dict) else {}
                    raw["property_extraction"] = {
                        "status": status,
                        "properties": [item for item in properties if isinstance(item, dict)],
                    }
                    result.raw_json = raw
                    # A completed independent lease decision is sufficient to
                    # make a previously OCR-misclassified contract eligible;
                    # non-lease rows are never promoted.
                    result.document_category = "lease_contract"
                    return result
                return result
            finally:
                workbook.close()

    def upsert(self, candidate: PdfCandidate, result: ExtractionResult) -> int:
        """Atomically replace every property row belonging to one contract."""
        rows = build_property_table_rows(result)
        processed_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        operational = {
            "processed_at": processed_at,
            "source_file_name": candidate.source_path.name,
            "copied_file_path": str(candidate.copied_path),
            "sha256": candidate.sha256,
            "file_created_at": candidate.created_at.isoformat(),
            "file_modified_at": candidate.modified_at.isoformat(),
        }
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            sha_col = PROPERTY_TABLE_COLUMNS.index("sha256") + 1
            for row_index in range(sheet.max_row, 1, -1):
                if str(sheet.cell(row=row_index, column=sha_col).value or "") == candidate.sha256:
                    sheet.delete_rows(row_index)
            for property_row in rows:
                payload = {**property_row, **operational}
                payload["property_row_id"] = (
                    f"{candidate.sha256}::{payload.get('property_row_id') or 'unresolved'}"
                )
                sheet.append([
                    _excel_value(payload.get(column, ""))
                    for column in PROPERTY_TABLE_COLUMNS
                ])
            self._format(sheet)
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
            workbook.close()
        return len(rows)

    def _load(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError as exc:
            raise RuntimeError("Missing dependency: install openpyxl with `pip install -r requirements.txt`.") from exc

        if self.path.exists():
            workbook = load_workbook(self.path)
            if PROPERTY_TABLE_SHEET_NAME in workbook.sheetnames:
                sheet = workbook[PROPERTY_TABLE_SHEET_NAME]
            else:
                sheet = workbook.create_sheet(PROPERTY_TABLE_SHEET_NAME)
                sheet.append(PROPERTY_TABLE_COLUMNS)
                _add_table(sheet, "PropertyTable", len(PROPERTY_TABLE_COLUMNS))
            migrated = _ensure_named_headers(sheet, PROPERTY_TABLE_COLUMNS)
            if not sheet.tables:
                _add_table(sheet, "PropertyTable", len(PROPERTY_TABLE_COLUMNS))
                migrated = True
            if migrated:
                self._format(sheet)
                workbook.save(self.path)
            return workbook, sheet

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = PROPERTY_TABLE_SHEET_NAME
        sheet.append(PROPERTY_TABLE_COLUMNS)
        _add_table(sheet, "PropertyTable", len(PROPERTY_TABLE_COLUMNS))
        self._format(sheet)
        return workbook, sheet

    def _format(self, sheet) -> None:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        widths = {
            "property_table_version": 16,
            "property_row_id": 78,
            "property_index": 14,
            "property_count": 14,
            "property_row_status": 24,
            "property_match_status": 30,
            "property_association_status": 24,
            "property_association_evidence": 48,
            "contract_operational_property_id": 28,
            "property_source_pages": 22,
            "property_structure_status": 24,
            "property_owner_status": 24,
            "property_source_file": 34,
            "property_total_area_m2": 22,
            "property_json": 72,
            "processed_at": 22,
            "source_file_name": 38,
            "copied_file_path": 48,
            "sha256": 66,
            "lessor": 38,
            "lessee": 34,
            "property_name": 42,
            "property_matrix_key": 20,
            "property_parish": 24,
            "property_municipality": 24,
            "property_district": 22,
            "property_total_area": 20,
            "owner_name": 38,
            "review_reason": 50,
            "raw_json": 72,
        }
        for index, header in enumerate(PROPERTY_TABLE_COLUMNS, start=1):
            if header in widths:
                sheet.column_dimensions[_column_letter(index)].width = widths[header]
        area_column = PROPERTY_TABLE_COLUMNS.index("property_total_area_m2") + 1
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row=row_index, column=area_column).number_format = '#,##0.00'
        if sheet.tables:
            table = next(iter(sheet.tables.values()))
            _set_table_ref(sheet, table, len(PROPERTY_TABLE_COLUMNS))


def _ensure_headers(sheet) -> bool:
    existing_headers = [
        sheet.cell(row=1, column=column).value
        for column in range(1, max(sheet.max_column, len(REGISTER_COLUMNS)) + 1)
    ]
    existing_headers = [str(header) if header is not None else "" for header in existing_headers]
    if existing_headers[: len(REGISTER_COLUMNS)] == REGISTER_COLUMNS:
        return False

    rows: list[dict[str, object]] = []
    for row_index in range(2, sheet.max_row + 1):
        row_data: dict[str, object] = {}
        for column_index, header in enumerate(existing_headers, start=1):
            if header:
                row_data[header] = sheet.cell(row=row_index, column=column_index).value
        rows.append(row_data)

    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    if sheet.max_column:
        sheet.delete_cols(1, sheet.max_column)
    for column_index, header in enumerate(REGISTER_COLUMNS, start=1):
        sheet.cell(row=1, column=column_index, value=header)
    for row_data in rows:
        sheet.append([row_data.get(column, "") for column in REGISTER_COLUMNS])
    return True


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


class PropertyExcelRegister:
    """Dedicated worksheet for the focused property-extraction pipeline."""

    def __init__(self, path: Path):
        self.path = path

    def existing_hashes(self, pipeline_version: str | None = None) -> set[str]:
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            sha_col = PROPERTY_REGISTER_COLUMNS.index("sha256") + 1
            version_col = PROPERTY_REGISTER_COLUMNS.index("pipeline_version") + 1
            values = {
                str(sheet.cell(row=row, column=sha_col).value)
                for row in range(2, sheet.max_row + 1)
                if sheet.cell(row=row, column=sha_col).value
                and (
                    pipeline_version is None
                    or sheet.cell(row=row, column=version_col).value == pipeline_version
                )
            }
            workbook.close()
            return values

    def append(self, payload: dict[str, object]) -> None:
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            sheet.append([
                _excel_value(payload.get(column, ""))
                for column in PROPERTY_REGISTER_COLUMNS
            ])
            self._format(sheet)
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
            workbook.close()

    def upsert(self, payload: dict[str, object]) -> None:
        """Replace older pipeline results for the same PDF hash."""
        sha256 = str(payload.get("sha256") or "")
        if not sha256:
            raise ValueError("Property result requires sha256.")
        with _workbook_lock(self.path):
            workbook, sheet = self._load()
            sha_col = PROPERTY_REGISTER_COLUMNS.index("sha256") + 1
            for row in range(sheet.max_row, 1, -1):
                if str(sheet.cell(row=row, column=sha_col).value or "") == sha256:
                    sheet.delete_rows(row)
            sheet.append([
                _excel_value(payload.get(column, ""))
                for column in PROPERTY_REGISTER_COLUMNS
            ])
            self._format(sheet)
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
            workbook.close()

    def _load(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError as exc:
            raise RuntimeError("Missing dependency: install openpyxl with `pip install -r requirements.txt`.") from exc

        if self.path.exists():
            workbook = load_workbook(self.path)
            if PROPERTY_SHEET_NAME in workbook.sheetnames:
                sheet = workbook[PROPERTY_SHEET_NAME]
            else:
                sheet = workbook.create_sheet(PROPERTY_SHEET_NAME)
                sheet.append(PROPERTY_REGISTER_COLUMNS)
                _add_table(sheet, "PropertyExtraction", len(PROPERTY_REGISTER_COLUMNS))
            migrated = _ensure_property_headers(sheet)
            if not sheet.tables:
                _add_table(sheet, "PropertyExtraction", len(PROPERTY_REGISTER_COLUMNS))
                migrated = True
            if migrated:
                self._format(sheet)
                workbook.save(self.path)
            return workbook, sheet

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = PROPERTY_SHEET_NAME
        sheet.append(PROPERTY_REGISTER_COLUMNS)
        _add_table(sheet, "PropertyExtraction", len(PROPERTY_REGISTER_COLUMNS))
        self._format(sheet)
        return workbook, sheet

    def _format(self, sheet) -> None:
        sheet.freeze_panes = "A2"
        widths = {
            "pipeline_version": 16, "processed_at": 22, "source_file_name": 34, "source_file_path": 48,
            "sha256": 66, "property_name": 42, "property_matrix_key": 18, "reason": 28,
            "extraction_notes": 54, "llm_error": 48, "evidence_model": 80,
            "property_catalog_text_source": 28, "property_pack_match_method": 34,
            "caderneta_source_file": 34, "crp_source_file": 34,
            "caderneta_evidence_sources": 46,
            "internal_caderneta_status": 38,
            "internal_caderneta_count": 18,
            "properties": 80,
            "audit": 80,
        }
        for index, header in enumerate(PROPERTY_REGISTER_COLUMNS, start=1):
            if header in widths:
                sheet.column_dimensions[_column_letter(index)].width = widths[header]
        if sheet.tables:
            table = next(iter(sheet.tables.values()))
            _set_table_ref(sheet, table, len(PROPERTY_REGISTER_COLUMNS))


def _excel_value(value: object) -> object:
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value


def _sheet_headers(sheet) -> dict[str, int]:
    if sheet.max_row < 1:
        return {}
    return {
        str(cell.value): index
        for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=False)), start=1)
        if cell.value
    }


def _sheet_row_dicts(sheet, headers: dict[str, int] | None = None):
    headers = headers or _sheet_headers(sheet)
    for values in sheet.iter_rows(min_row=2, values_only=True):
        yield {
            name: values[index - 1] if index <= len(values) else ""
            for name, index in headers.items()
        }


def _sheet_rows_by_sha(sheet) -> dict[str, dict[str, object]]:
    headers = _sheet_headers(sheet)
    if "sha256" not in headers:
        return {}
    return {
        str(row["sha256"]): row
        for row in _sheet_row_dicts(sheet, headers)
        if row.get("sha256")
    }


def _materialized_result(
    register_row: dict[str, object], extraction_row: dict[str, object]
) -> ExtractionResult:
    result = ExtractionResult()
    for field_name in vars(result):
        if field_name == "raw_json":
            continue
        value = register_row.get(field_name, "")
        if value not in (None, ""):
            setattr(result, field_name, value)
    result.document_category = "lease_contract"
    if not result.document_type:
        result.document_type = "lease_contract"
    raw = _json_dict(register_row.get("raw_json"))
    properties = _json_list(extraction_row.get("properties"))
    if not properties:
        properties = [_property_extraction_single_row(extraction_row)]
    raw["property_extraction"] = {
        "status": str(extraction_row.get("status") or ""),
        "properties": [item for item in properties if isinstance(item, dict)],
    }
    result.raw_json = raw
    return result


def _materialized_operational(
    register_row: dict[str, object], extraction_row: dict[str, object], sha256: str
) -> dict[str, object]:
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    return {
        "processed_at": now,
        "source_file_name": register_row.get("source_file_name") or extraction_row.get("source_file_name") or "",
        "copied_file_path": register_row.get("copied_file_path") or extraction_row.get("source_file_path") or "",
        "sha256": sha256,
        "file_created_at": register_row.get("file_created_at") or "",
        "file_modified_at": register_row.get("file_modified_at") or "",
    }


def _json_dict(value: object) -> dict[str, object]:
    if isinstance(value, dict):
        return dict(value)
    if not isinstance(value, str) or not value.strip():
        return {}
    try:
        parsed = json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return {}
    return dict(parsed) if isinstance(parsed, dict) else {}


def _json_list(value: object) -> list[dict[str, object]]:
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    if not isinstance(value, str) or not value.strip():
        return []
    try:
        parsed = json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return []
    return [item for item in parsed if isinstance(item, dict)] if isinstance(parsed, list) else []


def _property_extraction_single(
    values: tuple[object, ...], headers: dict[str, int]
) -> dict[str, object]:
    """Convert a legacy single-property Property Extraction row to a list item."""
    def value(name: str) -> object:
        index = headers.get(name)
        return values[index - 1] if index and index <= len(values) else ""

    return {
        "property_name": value("property_name"),
        "matrix_article": value("matrix_article"),
        "matrix_section": value("matrix_section"),
        "property_matrix_key": value("property_matrix_key"),
        "area_m2": value("area_m2"),
        "owner_name": value("owner_name"),
        "source_file": value("source_file_name"),
    }


def _property_extraction_single_row(row: dict[str, object]) -> dict[str, object]:
    return {
        "property_name": row.get("property_name", ""),
        "matrix_article": row.get("matrix_article", ""),
        "matrix_section": row.get("matrix_section", ""),
        "property_matrix_key": row.get("property_matrix_key", ""),
        "area_m2": row.get("area_m2", ""),
        "owner_name": row.get("owner_name", ""),
        "source_file": row.get("source_file_name", ""),
    }


def _ensure_named_headers(sheet, columns: tuple[str, ...]) -> bool:
    existing_headers = [
        str(sheet.cell(row=1, column=column).value or "")
        for column in range(1, max(sheet.max_column, len(columns)) + 1)
    ]
    if existing_headers[:len(columns)] == list(columns):
        return False

    rows = [
        {
            header: sheet.cell(row=row_index, column=column_index).value
            for column_index, header in enumerate(existing_headers, start=1)
            if header
        }
        for row_index in range(2, sheet.max_row + 1)
    ]
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    if sheet.max_column:
        sheet.delete_cols(1, sheet.max_column)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    return True


def _ensure_property_headers(sheet) -> bool:
    existing_headers = [
        str(sheet.cell(row=1, column=column).value or "")
        for column in range(1, max(sheet.max_column, len(PROPERTY_REGISTER_COLUMNS)) + 1)
    ]
    if existing_headers[:len(PROPERTY_REGISTER_COLUMNS)] == list(PROPERTY_REGISTER_COLUMNS):
        return False

    rows: list[dict[str, object]] = []
    for row_index in range(2, sheet.max_row + 1):
        rows.append({
            header: sheet.cell(row=row_index, column=column_index).value
            for column_index, header in enumerate(existing_headers, start=1)
            if header
        })
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    if sheet.max_column:
        sheet.delete_cols(1, sheet.max_column)
    for column_index, header in enumerate(PROPERTY_REGISTER_COLUMNS, start=1):
        sheet.cell(row=1, column=column_index, value=header)
    for row in rows:
        sheet.append([row.get(column, "") for column in PROPERTY_REGISTER_COLUMNS])
    return True


def _add_table(sheet, name: str, column_count: int) -> None:
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    table = Table(displayName=name, ref=f"A1:{_column_letter(column_count)}1")
    table.autoFilter = AutoFilter(ref=table.ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _set_table_ref(sheet, table, column_count: int) -> None:
    """Keep the table and its serialized AutoFilter range in lockstep.

    ``openpyxl`` does not update ``table.autoFilter.ref`` when ``table.ref``
    changes. Excel repairs a workbook when those ranges diverge, sometimes by
    deleting the complete table. Every append, replacement and rebuild uses
    this helper after changing a sheet's rows.
    """
    table.ref = f"A1:{_column_letter(column_count)}{max(sheet.max_row, 1)}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref


@contextmanager
def _workbook_lock(path: Path):
    """Serialize local workbook writes when both pipelines run together."""
    lock_path = path.with_suffix(path.suffix + ".lock")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    with lock_path.open("a+") as handle:
        if fcntl:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            if fcntl:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
