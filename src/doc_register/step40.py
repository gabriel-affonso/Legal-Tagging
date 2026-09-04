"""Step 4.0 normalized workbook output.

The document register remains the forensic record.  This module is the only
place that turns its consolidated evidence into operational entities.  In
particular, writers must not bypass :class:`NormalizedWorkbookRegister` to
write a ``db.*`` worksheet directly.
"""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
import logging
from pathlib import Path
import re
import shutil
import unicodedata
from typing import Any, Iterable, Mapping

from .models import ExtractionResult, PdfCandidate
from .schemas import REGISTER_COLUMNS


LOGGER = logging.getLogger(__name__)
NORMALIZED_WORKBOOK_SCHEMA_VERSION = "3"

PROCESSING_PENDING = "PENDENTE"
PROCESSING_DONE = "PROCESSADA"
PROCESSING_PARTIAL = "PROCESSADA_PARCIAL"
PROCESSING_REVIEW = "REQUER_REVISAO"

ENTRY_COLUMNS = (
    "entrada_id", "referencia_mg*", "codigo_interno", "nome_proprietario", "nif",
    "documento_id", "tipo_proprietario", "papel_no_contrato", "tipo_direito_terreno",
    "quota", "telefone", "email", "morada", "iban", "distrito", "concelho",
    "freguesia", "local", "artigo_matricial", "parte", "area_at_m2",
    "descricao_predial", "area_medida_m2", "area_contratada_ha", "tipo_contrato",
    "lote", "referencia_contrato", "data_assinatura", "projeto", "transferencia_newco",
    "observacoes", "estado_processamento", "mensagem_processamento", "processado_em",
    "contrato_id", "proprietario_id", "terreno_id",
)

TABLE_SCHEMAS: dict[str, tuple[str, ...]] = {
    "db.Contrato": (
        "contrato_id", "referencia_mg", "lote", "referencia_contrato", "tipo",
        "data_assinatura", "projeto", "transferencia_newco", "observacoes", "origem",
        "atualizado_em", "area_documental_total_ha", "status_area_documental",
        "area_parcelas_considerada_ha", "diferenca_documental_ha",
    ),
    "db.Terreno": (
        "parcela_id", "terreno_grupo_id", "codigo_interno", "distrito", "concelho",
        "freguesia", "local", "artigo_matricial", "parte", "area_at_m2", "area_at_ha",
        "descricao_predial", "area_medida_m2", "area_medida_ha", "area_registada_original_ha",
        "area_documental_parcela_ha", "area_considerada_ha", "fonte_area_considerada",
        "status_validacao_area", "observacoes", "origem", "linha_origem",
    ),
    "db.Proprietario": (
        "proprietario_id", "nome_canonico", "tipo", "documento_id", "nif", "morada",
        "estado", "origem", "observacoes",
    ),
    "db.Contacto": (
        "contacto_id", "proprietario_id", "tipo", "valor", "principal", "estado",
        "origem", "observacoes",
    ),
    "db.ContaBancaria": ("conta_id", "proprietario_id", "iban", "origem", "observacoes"),
    "db.ContratoProprietario": (
        "relacao_id", "contrato_id", "proprietario_id", "papel", "quota", "origem", "observacoes",
    ),
    "db.ContratoTerreno": (
        "relacao_id", "contrato_id", "parcela_id", "terreno_grupo_id_legado",
        "area_registada_original_ha", "area_documental_parcela_ha", "area_considerada_ha",
        "fonte_area_considerada", "status_validacao_area", "origem", "linha_origem",
    ),
    # ``terreno_id`` is deliberately the physical ``parcela_id``.  This is
    # kept because the operational template has that legacy column name.
    "db.ProprietarioTerreno": (
        "relacao_id", "proprietario_id", "terreno_id", "tipo_direito", "quota", "origem", "observacoes",
    ),
    "db.Processamento": (
        "entrada_id", "estado", "mensagem", "processado_em", "contrato_id",
        "proprietario_id", "terreno_id",
    ),
}

AUDIT_SCHEMAS: dict[str, tuple[str, ...]] = {
    "audit.DocumentRegister": REGISTER_COLUMNS,
    "audit.PropertyEvidence": (
        "source_sha256", "property_index", "parcela_candidate_key", "values_json",
        "source_pages", "authority", "association_status",
    ),
    "audit.FieldEvidence": (
        "source_sha256", "entity_table", "entity_id", "field", "value", "authority", "evidence_json",
    ),
    "audit.Conflicts": (
        "source_sha256", "entry_id", "entity_table", "entity_id", "field", "message", "candidates_json",
    ),
}

TABLE_KEYS = {
    "db.Contrato": ("contrato_id",), "db.Terreno": ("parcela_id",),
    "db.Proprietario": ("proprietario_id",), "db.Contacto": ("contacto_id",),
    "db.ContaBancaria": ("conta_id",), "db.ContratoProprietario": ("relacao_id",),
    "db.ContratoTerreno": ("relacao_id",), "db.ProprietarioTerreno": ("relacao_id",),
    "db.Processamento": ("entrada_id",),
}

# The available reference workbook predates schema V3.  These mappings retain
# its data during a one-time column migration instead of silently discarding
# it when the physical parcel schema is introduced.
LEGACY_COLUMN_MAP = {
    "db.Terreno": {
        "terreno_id": "parcela_id", "area_contratada_ha": "area_documental_parcela_ha",
    },
    "db.ContratoTerreno": {
        "terreno_id": "parcela_id", "area_contratada_ha": "area_documental_parcela_ha",
    },
}

CONTROLLED_DEFAULTS = {
    "Tipo_Proprietario": ("PESSOA", "EMPRESA", "HERANCA", "ENTIDADE_PUBLICA", "OUTRA_ENTIDADE"),
    "Papel_Contrato": ("PROPRIETARIO", "COPROPRIETARIO", "USUFRUTUARIO", "REPRESENTANTE", "CABECA_DE_CASAL", "PROCURADOR", "SIGNATARIO", "OUTRO"),
    "Tipo_Direito_Terreno": ("PROPRIEDADE", "COPROPRIEDADE", "USUFRUTO", "NUA_PROPRIEDADE", "POSSE", "REPRESENTACAO", "OUTRO"),
    "Transferencia_NewCo": ("Yes", "No"),
}

TEXT_COLUMNS = {
    "entrada_id", "referencia_mg*", "codigo_interno", "nif", "documento_id", "iban",
    "contrato_id", "proprietario_id", "terreno_id", "parcela_id", "terreno_grupo_id",
    "terreno_grupo_id_legado", "relacao_id", "contacto_id", "conta_id", "artigo_matricial",
}
DATE_COLUMNS = {"data_assinatura"}
DATETIME_COLUMNS = {"processado_em", "atualizado_em"}


def text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def comparison_key(value: object) -> str:
    folded = "".join(
        char for char in unicodedata.normalize("NFKD", text(value).casefold())
        if not unicodedata.combining(char)
    )
    return re.sub(r"[^a-z0-9]+", "", folded)


def stable_id(prefix: str, *parts: object) -> str:
    payload = "|".join(comparison_key(part) for part in parts)
    return f"{prefix}-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:16].upper()}"


def normalise_nif(value: object) -> str:
    return re.sub(r"\D", "", text(value))


def valid_nif(value: object) -> bool:
    nif = normalise_nif(value)
    if len(nif) != 9 or nif == "0" * 9:
        return False
    total = sum(int(digit) * weight for digit, weight in zip(nif[:8], range(9, 1, -1)))
    check = 11 - total % 11
    return int(nif[-1]) == (0 if check >= 10 else check)


def normalise_iban(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", text(value).upper())


def valid_iban(value: object) -> bool:
    iban = normalise_iban(value)
    if len(iban) < 15 or not re.fullmatch(r"[A-Z]{2}[0-9]{2}[A-Z0-9]+", iban):
        return False
    rearranged = iban[4:] + iban[:4]
    numeric = "".join(str(ord(char) - 55) if char.isalpha() else char for char in rearranged)
    return int(numeric) % 97 == 1


def as_decimal(value: object) -> Decimal | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return Decimal(str(value))
    candidate = text(value).replace(" ", "")
    if "," in candidate and "." in candidate:
        candidate = candidate.replace(".", "").replace(",", ".")
    else:
        candidate = candidate.replace(",", ".")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", candidate)
    if not match:
        return None
    try:
        return Decimal(match.group(0))
    except InvalidOperation:
        return None


def parse_area(value: object) -> tuple[Decimal | None, Decimal | None]:
    """Return (m2, ha), accepting explicit units and never guessing them."""
    numeric = as_decimal(value)
    if numeric is None:
        return None, None
    raw = text(value).casefold().replace("²", "2")
    if re.search(r"\b(ha|hectare|hectares)\b", raw):
        return numeric * Decimal("10000"), numeric
    if re.search(r"\b(m2|metro|metros)\b", raw):
        return numeric, numeric / Decimal("10000")
    # The staging columns themselves define their units.
    return numeric, numeric / Decimal("10000")


def _canonical_choice(value: object, choices: Iterable[str], aliases: Mapping[str, str], *, fallback: str = "") -> tuple[str, str]:
    raw = text(value)
    if not raw:
        return fallback, ""
    key = comparison_key(raw)
    mapped = aliases.get(key, raw.upper())
    allowed = {comparison_key(item): item for item in choices}
    return (allowed[comparison_key(mapped)], "") if comparison_key(mapped) in allowed else (fallback, f"valor não reconhecido: {raw}")


def _owner_type(value: object) -> tuple[str, str]:
    return _canonical_choice(value, CONTROLLED_DEFAULTS["Tipo_Proprietario"], {
        "pessoa": "PESSOA", "pessoasingular": "PESSOA", "particular": "PESSOA",
        "sociedade": "EMPRESA", "pessoacoletiva": "EMPRESA", "empresa": "EMPRESA",
        "heranca": "HERANCA", "entidadepublica": "ENTIDADE_PUBLICA",
    }, fallback="PESSOA")


def _contract_role(value: object) -> tuple[str, str]:
    return _canonical_choice(value, CONTROLLED_DEFAULTS["Papel_Contrato"], {
        "senhorio": "PROPRIETARIO", "proprietario": "PROPRIETARIO", "coproprietario": "COPROPRIETARIO",
    }, fallback="PROPRIETARIO")


def _right_type(value: object) -> tuple[str, str]:
    return _canonical_choice(value, CONTROLLED_DEFAULTS["Tipo_Direito_Terreno"], {
        "propriedade": "PROPRIEDADE", "copropriedade": "COPROPRIEDADE", "usufruto": "USUFRUTO",
        "nuapropriedade": "NUA_PROPRIEDADE", "posse": "POSSE",
    }, fallback="PROPRIEDADE")


def _reference_from(candidate: PdfCandidate, context: str = "", metadata: object = "", document_text: str = "") -> tuple[str, str]:
    if text(metadata):
        return text(metadata), "metadado_estruturado"
    if text(context):
        return text(context), "contexto_operacional"
    # A filename is accepted only for an explicitly-labelled MG reference, not
    # as a substitute for an arbitrary contract identity.
    pattern = r"\bMG[\s_-]*[A-Z0-9][A-Z0-9_-]*\b"
    path_match = re.search(pattern, str(candidate.source_path.parent).upper())
    if path_match:
        return path_match.group(0).replace("_", "-").replace(" ", "-"), "contexto_pasta"
    filename_match = re.search(pattern, candidate.source_path.name.upper())
    if filename_match:
        return filename_match.group(0).replace("_", "-").replace(" ", "-"), "nome_ficheiro"
    text_match = re.search(pattern, document_text.upper())
    return (text_match.group(0).replace("_", "-").replace(" ", "-"), "texto_documento") if text_match else ("", "")


def _mapping_list(value: object) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [dict(item) for item in value if isinstance(item, Mapping)]
    if isinstance(value, tuple):
        return [dict(item) for item in value if isinstance(item, Mapping)]
    return []


def _properties(result: ExtractionResult) -> list[dict[str, Any]]:
    raw = result.raw_json if isinstance(result.raw_json, Mapping) else {}
    candidates: list[object] = []
    for parent, key in (
        ("final_contract_resolution", "properties"), ("step2_7_final_resolution", "properties"),
        ("property_extraction", "properties"), ("step3_3_field_centric", "cadastral_properties"),
    ):
        holder = raw.get(parent)
        candidates.append(holder.get(key) if isinstance(holder, Mapping) else None)
    candidates.append(raw.get("internal_caderneta_groups"))
    for candidate in candidates:
        values = _mapping_list(candidate)
        if values:
            return values
    scalar = {
        "codigo_interno": result.property_number, "property_name": result.property_display_name or result.property_name,
        "property_article": result.property_article, "property_section": result.property_section,
        "property_parish": result.property_parish, "property_municipality": result.property_municipality,
        "property_district": result.property_district, "property_location": result.property_location,
        "property_total_area": result.property_total_area, "leased_parcel_area": result.leased_parcel_area,
        "descricao_predial": result.registry_description, "owner_name": result.owner_name,
        "owner_tax_id": result.owner_tax_id, "owner_address": result.owner_address,
    }
    return [scalar] if any(text(value) for value in scalar.values()) else []


def _owner_candidates(result: ExtractionResult) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for attribute, role in (
        ("cadastral_owners", "PROPRIETARIO"), ("registered_owners", "PROPRIETARIO"),
        ("owners", "PROPRIETARIO"), ("contract_lessor", "PROPRIETARIO"),
    ):
        for owner in _mapping_list(getattr(result, attribute, ())):
            candidates.append({**owner, "papel": owner.get("papel") or owner.get("role") or role})
    if not candidates and text(result.lessor):
        candidates.append({"name": result.lessor, "nif": result.lessor_tax_id, "papel": "PROPRIETARIO"})
    if not candidates and text(result.owner_name):
        candidates.append({"name": result.owner_name, "nif": result.owner_tax_id, "address": result.owner_address, "papel": "PROPRIETARIO"})
    return candidates


def _owner_name(owner: Mapping[str, Any]) -> str:
    return text(owner.get("nome") or owner.get("name") or owner.get("owner_name") or owner.get("nome_canonico"))


def _owner_nif(owner: Mapping[str, Any]) -> str:
    return normalise_nif(owner.get("nif") or owner.get("tax_id") or owner.get("owner_tax_id"))


def _owner_matches(left: Mapping[str, Any], right: Mapping[str, Any]) -> bool:
    left_nif, right_nif = _owner_nif(left), _owner_nif(right)
    return bool(left_nif and right_nif and left_nif == right_nif) or (
        bool(_owner_name(left)) and comparison_key(_owner_name(left)) == comparison_key(_owner_name(right))
    )


def build_staging_rows(candidate: PdfCandidate, result: ExtractionResult, *, reference_context: str = "", batch_id: str = "", document_text: str = "") -> list[dict[str, Any]]:
    """Consolidate one document into safe operational combinations.

    Property holders generate owner-parcel rows.  Contract parties without a
    documented parcel generate party-only rows.  This is intentionally not a
    cartesian product.
    """
    raw = result.raw_json if isinstance(result.raw_json, Mapping) else {}
    metadata_reference = raw.get("referencia_mg") or raw.get("reference_mg")
    reference, reference_source = _reference_from(candidate, reference_context, metadata_reference, document_text)
    properties = _properties(result)
    owners = _owner_candidates(result)
    rows: list[dict[str, Any]] = []

    def staging(owner: Mapping[str, Any] | None, prop: Mapping[str, Any] | None, index: int, *, property_link: bool) -> dict[str, Any]:
        owner = owner or {}
        prop = prop or {}
        owner_name = _owner_name(owner) or text(prop.get("owner_name"))
        owner_nif = _owner_nif(owner) or normalise_nif(prop.get("owner_tax_id"))
        article = text(prop.get("property_article") or prop.get("matrix_article") or prop.get("artigo_matricial"))
        area_m2, _ = parse_area(prop.get("property_total_area_m2") or prop.get("area_m2") or prop.get("property_total_area") or prop.get("area_at_m2"))
        measured_m2, _ = parse_area(prop.get("area_medida_m2"))
        contractual = as_decimal(prop.get("leased_parcel_area") or prop.get("area_contratada_ha"))
        if contractual is not None and re.search(r"\b(m2|m²)\b", text(prop.get("leased_parcel_area") or prop.get("area_contratada_ha")).casefold()):
            contractual /= Decimal("10000")
        entry = stable_id("ENT", candidate.sha256, reference, owner_nif or owner_name, text(prop.get("codigo_interno") or prop.get("property_number")), article, index)
        role, role_problem = _contract_role(owner.get("papel") or owner.get("role"))
        owner_type, owner_type_problem = _owner_type(owner.get("tipo") or owner.get("type"))
        right, right_problem = _right_type(owner.get("tipo_direito") or owner.get("right_type"))
        observations = [part for part in (text(result.extraction_notes), f"referencia_mg: {reference_source}" if reference_source else "", role_problem, owner_type_problem, right_problem) if part]
        return {
            "entrada_id": entry, "referencia_mg*": reference, "codigo_interno": text(prop.get("codigo_interno") or prop.get("property_number")),
            "nome_proprietario": owner_name, "nif": owner_nif, "documento_id": text(owner.get("documento_id") or owner.get("document_id")),
            "tipo_proprietario": owner_type, "papel_no_contrato": role, "tipo_direito_terreno": right,
            "quota": text(owner.get("quota") or owner.get("share")), "telefone": text(owner.get("telefone") or owner.get("phone")),
            "email": text(owner.get("email")), "morada": text(owner.get("morada") or owner.get("address") or prop.get("owner_address")),
            "iban": "", "distrito": text(prop.get("property_district") or prop.get("distrito")),
            "concelho": text(prop.get("property_municipality") or prop.get("concelho")), "freguesia": text(prop.get("property_parish") or prop.get("freguesia")),
            "local": text(prop.get("property_location") or prop.get("local")), "artigo_matricial": article,
            "parte": text(prop.get("property_section") or prop.get("matrix_section") or prop.get("parte")),
            "area_at_m2": area_m2, "descricao_predial": text(prop.get("descricao_predial") or prop.get("registry_description") or prop.get("property_name")),
            "area_medida_m2": measured_m2, "area_contratada_ha": contractual,
            "tipo_contrato": text(result.contract_type), "lote": text(batch_id), "referencia_contrato": text(result.property_number),
            "data_assinatura": text(result.signed_date), "projeto": "", "transferencia_newco": "", "observacoes": " | ".join(observations),
            "estado_processamento": PROCESSING_PENDING, "mensagem_processamento": "", "processado_em": "",
            "contrato_id": "", "proprietario_id": "", "terreno_id": "", "_property_link": property_link,
        }

    used_owner_keys: set[str] = set()
    for property_index, prop in enumerate(properties, start=1):
        property_owner = {"name": prop.get("owner_name"), "nif": prop.get("owner_tax_id"), "address": prop.get("owner_address")}
        if _owner_name(property_owner):
            matching = next((item for item in owners if _owner_matches(item, property_owner)), property_owner)
            rows.append(staging(matching, prop, property_index, property_link=True))
            used_owner_keys.add(_owner_nif(matching) or comparison_key(_owner_name(matching)))
        else:
            rows.append(staging(None, prop, property_index, property_link=False))
    for owner_index, owner in enumerate(owners, start=len(rows) + 1):
        key = _owner_nif(owner) or comparison_key(_owner_name(owner))
        if key and key in used_owner_keys:
            continue
        rows.append(staging(owner, None, owner_index, property_link=False))
    if not rows:
        rows.append(staging(None, None, 1, property_link=False))

    # Bank information is operational only when a document holder matches a
    # row owner.  Otherwise the value stays in the audit record.
    holder = {"name": result.bank_account_holder}
    for row in rows:
        if text(result.iban) and text(row["nome_proprietario"]) and _owner_matches(holder, {"name": row["nome_proprietario"]}):
            row["iban"] = normalise_iban(result.iban)
    return rows


class NormalizedWorkbookRegister:
    """Versioned, idempotent normalized workbook exporter and normalizer."""

    def __init__(self, path: Path, *, template_path: Path | None = None):
        self.path = path
        self.template_path = template_path

    def existing_hashes(self) -> set[str]:
        if not self.path.exists():
            return set()
        from openpyxl import load_workbook
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            if "audit.DocumentRegister" not in workbook.sheetnames:
                return set()
            sheet = workbook["audit.DocumentRegister"]
            headers = _headers(sheet)
            column = headers.get("sha256")
            return {text(sheet.cell(row, column).value) for row in range(2, sheet.max_row + 1) if column and text(sheet.cell(row, column).value)}
        finally:
            workbook.close()

    def publish(self, candidate: PdfCandidate, result: ExtractionResult, *, reference_context: str = "", batch_id: str = "", document_text: str = "") -> dict[str, int]:
        workbook = self._load()
        try:
            self._audit_document(workbook, candidate, result)
            rows = build_staging_rows(candidate, result, reference_context=reference_context, batch_id=batch_id, document_text=document_text)
            self._audit_properties(workbook, candidate, result, rows)
            self._upsert_staging(workbook["Entrada_Rapida"], rows)
            counts = self._normalise(workbook)
            workbook.properties.keywords = f"normalized-workbook-schema={NORMALIZED_WORKBOOK_SCHEMA_VERSION}"
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
            LOGGER.info("Step 4.0 workbook updated: entries=%s processed=%s partial=%s review=%s", len(rows), counts[PROCESSING_DONE], counts[PROCESSING_PARTIAL], counts[PROCESSING_REVIEW])
            return counts
        finally:
            workbook.close()

    def process_pending(self) -> dict[str, int]:
        workbook = self._load()
        try:
            counts = self._normalise(workbook)
            workbook.save(self.path)
            return counts
        finally:
            workbook.close()

    def _load(self):
        from openpyxl import Workbook, load_workbook
        if not self.path.exists() and self.template_path and self.template_path.exists() and self.template_path != self.path:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self.template_path, self.path)
        workbook = load_workbook(self.path) if self.path.exists() else Workbook()
        if workbook.active.title == "Sheet" and len(workbook.sheetnames) == 1:
            workbook.active.title = "INSTRUCOES"
        self._ensure_schema(workbook)
        return workbook

    def _ensure_schema(self, workbook) -> None:
        required = ("INSTRUCOES", "Entrada_Rapida", *TABLE_SCHEMAS.keys(), "Listas", *AUDIT_SCHEMAS.keys())
        for name in required:
            if name not in workbook.sheetnames:
                workbook.create_sheet(name)
        for name, columns in {"Entrada_Rapida": ENTRY_COLUMNS, **TABLE_SCHEMAS, **AUDIT_SCHEMAS}.items():
            _ensure_sheet_columns(workbook[name], columns, _table_name(name))
        self._ensure_lists(workbook["Listas"])
        # Keep required operational tabs at the beginning. Other template tabs
        # (including legacy compatibility sheets) retain their contents.
        ordering = ["INSTRUCOES", "Entrada_Rapida", *TABLE_SCHEMAS.keys(), "Listas"]
        for offset, name in enumerate(ordering):
            sheet = workbook[name]
            workbook._sheets.remove(sheet)
            workbook._sheets.insert(offset, sheet)

    def _ensure_lists(self, sheet) -> None:
        headers = tuple(CONTROLLED_DEFAULTS)
        existing = tuple(text(cell.value) for cell in sheet[1]) if sheet.max_row else ()
        if existing[:len(headers)] != headers:
            if sheet.max_row:
                sheet.delete_rows(1, sheet.max_row)
            sheet.append(headers)
        for index, (_, values) in enumerate(CONTROLLED_DEFAULTS.items(), start=1):
            for row, value in enumerate(values, start=2):
                if not text(sheet.cell(row, index).value):
                    sheet.cell(row, index, value)

    def _audit_document(self, workbook, candidate: PdfCandidate, result: ExtractionResult) -> None:
        values = {
            "processed_at": datetime.now(timezone.utc).replace(microsecond=0),
            "source_file_name": candidate.source_path.name, "copied_file_path": str(candidate.copied_path),
            "sha256": candidate.sha256, "file_created_at": candidate.created_at,
            "file_modified_at": candidate.modified_at,
        }
        values.update({field: getattr(result, field, "") for field in REGISTER_COLUMNS})
        _upsert(workbook["audit.DocumentRegister"], AUDIT_SCHEMAS["audit.DocumentRegister"], ("sha256",), values)

    def _audit_properties(self, workbook, candidate: PdfCandidate, result: ExtractionResult, rows: list[dict[str, Any]]) -> None:
        sheet = workbook["audit.PropertyEvidence"]
        for index, property_value in enumerate(_properties(result), start=1):
            parcel_key = stable_id("PAR", property_value.get("codigo_interno") or property_value.get("property_number"), property_value.get("property_parish"), property_value.get("property_article") or property_value.get("matrix_article"), property_value.get("property_section") or property_value.get("matrix_section"))
            values = {
                "source_sha256": candidate.sha256, "property_index": index, "parcela_candidate_key": parcel_key,
                "values_json": property_value, "source_pages": property_value.get("pages") or property_value.get("page_numbers") or [],
                "authority": _authority_for(result.document_category, "property"),
                "association_status": property_value.get("association_status") or "document_context",
            }
            _upsert(sheet, AUDIT_SCHEMAS["audit.PropertyEvidence"], ("source_sha256", "property_index"), values)
            for field, value in property_value.items():
                if value in (None, "", [], {}):
                    continue
                _upsert(workbook["audit.FieldEvidence"], AUDIT_SCHEMAS["audit.FieldEvidence"], ("source_sha256", "entity_table", "entity_id", "field"), {
                    "source_sha256": candidate.sha256, "entity_table": "candidate.Terreno",
                    "entity_id": parcel_key, "field": field, "value": value,
                    "authority": _authority_for(result.document_category, "property"),
                    "evidence_json": {"pages": property_value.get("pages") or property_value.get("page_numbers") or []},
                })
        if text(result.cadastral_conflicts) or text(result.evidence_conflicts):
            self._audit_conflict(
                workbook, candidate.sha256, "", "audit.DocumentRegister", candidate.sha256,
                "cadastral_evidence", text(result.cadastral_conflicts) or text(result.evidence_conflicts),
            )
        for row in rows:
            if text(row.get("referencia_mg*")):
                continue
            self._audit_conflict(workbook, candidate.sha256, row["entrada_id"], "", "", "referencia_mg", "referencia_mg ausente")

    def _audit_conflict(self, workbook, source_hash: str, entry_id: str, table: str, entity_id: str, field: str, message: str) -> None:
        _upsert(workbook["audit.Conflicts"], AUDIT_SCHEMAS["audit.Conflicts"], ("source_sha256", "entry_id", "field", "message"), {
            "source_sha256": source_hash, "entry_id": entry_id, "entity_table": table,
            "entity_id": entity_id, "field": field, "message": message, "candidates_json": "",
        })

    def _upsert_staging(self, sheet, rows: list[dict[str, Any]]) -> None:
        for row in rows:
            output = {key: value for key, value in row.items() if not key.startswith("_")}
            existing = _row_by_key(sheet, ENTRY_COLUMNS, ("entrada_id",), (output["entrada_id"],))
            comparable = tuple(
                comparison_key(output.get(column))
                for column in ENTRY_COLUMNS
                if column not in {"estado_processamento", "mensagem_processamento", "processado_em", "contrato_id", "proprietario_id", "terreno_id"}
            )
            current = tuple(
                comparison_key(existing.get(column))
                for column in ENTRY_COLUMNS
                if column not in {"estado_processamento", "mensagem_processamento", "processado_em", "contrato_id", "proprietario_id", "terreno_id"}
            ) if existing else ()
            if existing and current == comparable:
                continue
            _upsert(sheet, ENTRY_COLUMNS, ("entrada_id",), output, replace=True)

    def _normalise(self, workbook) -> dict[str, int]:
        entries = workbook["Entrada_Rapida"]
        entry_headers = _headers(entries)
        counts = {PROCESSING_DONE: 0, PROCESSING_PARTIAL: 0, PROCESSING_REVIEW: 0}
        for row_number in range(2, entries.max_row + 1):
            entry = {name: entries.cell(row_number, column).value for name, column in entry_headers.items()}
            if not any(text(value) for value in entry.values()):
                continue
            if text(entry.get("estado_processamento")).upper() == PROCESSING_DONE:
                continue
            status, message, ids = self._normalise_entry(workbook, entry, row_number)
            for field, value in (("estado_processamento", status), ("mensagem_processamento", message), ("processado_em", datetime.now(timezone.utc).replace(microsecond=0)), *ids.items()):
                entries.cell(row_number, entry_headers[field], _excel_value(value, field))
            _upsert(workbook["db.Processamento"], TABLE_SCHEMAS["db.Processamento"], ("entrada_id",), {
                "entrada_id": entry["entrada_id"], "estado": status, "mensagem": message,
                "processado_em": datetime.now(timezone.utc).replace(microsecond=0), **ids,
            }, replace=True)
            counts[status] += 1
        self._reconcile_contract_areas(workbook)
        return counts

    def _normalise_entry(self, workbook, entry: Mapping[str, Any], source_row: int) -> tuple[str, str, dict[str, str]]:
        reference = text(entry.get("referencia_mg*"))
        ids = {"contrato_id": "", "proprietario_id": "", "terreno_id": ""}
        if not reference:
            return PROCESSING_REVIEW, "referencia_mg ausente", ids
        contract_id = stable_id("CTR", reference)
        ids["contrato_id"] = contract_id
        _upsert(workbook["db.Contrato"], TABLE_SCHEMAS["db.Contrato"], ("contrato_id",), {
            "contrato_id": contract_id, "referencia_mg": reference, "lote": entry.get("lote"),
            "referencia_contrato": entry.get("referencia_contrato"), "tipo": entry.get("tipo_contrato"),
            "data_assinatura": _as_date(entry.get("data_assinatura")), "projeto": entry.get("projeto"),
            "transferencia_newco": entry.get("transferencia_newco"), "observacoes": entry.get("observacoes"),
            "origem": "STEP4.0_ENTRADA_RAPIDA", "atualizado_em": datetime.now(timezone.utc).replace(microsecond=0),
        })
        messages: list[str] = []
        owner_id = self._owner(workbook, entry, reference)
        if text(entry.get("nif")) and not valid_nif(entry.get("nif")):
            messages.append("NIF inválido; não usado como identidade forte")
        _, owner_type_problem = _owner_type(entry.get("tipo_proprietario"))
        if owner_type_problem:
            messages.append(owner_type_problem)
        if owner_id:
            ids["proprietario_id"] = owner_id
            role, role_problem = _contract_role(entry.get("papel_no_contrato"))
            if role_problem:
                messages.append(role_problem)
            _upsert(workbook["db.ContratoProprietario"], TABLE_SCHEMAS["db.ContratoProprietario"], ("relacao_id",), {
                "relacao_id": stable_id("REL", "contract-owner", contract_id, owner_id), "contrato_id": contract_id,
                "proprietario_id": owner_id, "papel": role, "quota": text(entry.get("quota")),
                "origem": "STEP4.0_ENTRADA_RAPIDA", "observacoes": entry.get("observacoes"),
            })
            self._contacts_and_bank(workbook, entry, owner_id, messages)
        elif any(text(entry.get(field)) for field in ("nif", "documento_id", "telefone", "email", "morada", "iban", "papel_no_contrato", "quota")):
            messages.append("dados de proprietário ignorados porque falta nome_proprietario")
        else:
            messages.append("sem proprietário; relação contrato-proprietário não criada")
        parcel_id = self._parcel(workbook, entry, source_row)
        if parcel_id:
            ids["terreno_id"] = parcel_id
            terrain = _row_by_key(workbook["db.Terreno"], TABLE_SCHEMAS["db.Terreno"], ("parcela_id",), (parcel_id,))
            _upsert(workbook["db.ContratoTerreno"], TABLE_SCHEMAS["db.ContratoTerreno"], ("relacao_id",), {
                "relacao_id": stable_id("REL", "contract-parcel", contract_id, parcel_id), "contrato_id": contract_id,
                "parcela_id": parcel_id, "terreno_grupo_id_legado": terrain.get("terreno_grupo_id") if terrain else "",
                "area_registada_original_ha": terrain.get("area_registada_original_ha") if terrain else "",
                "area_documental_parcela_ha": entry.get("area_contratada_ha"),
                "area_considerada_ha": terrain.get("area_considerada_ha") if terrain else "",
                "fonte_area_considerada": terrain.get("fonte_area_considerada") if terrain else "",
                "status_validacao_area": terrain.get("status_validacao_area") if terrain else "",
                "origem": "STEP4.0_ENTRADA_RAPIDA", "linha_origem": source_row,
            })
        elif any(text(entry.get(field)) for field in ("codigo_interno", "distrito", "concelho", "freguesia", "local", "artigo_matricial", "parte", "area_at_m2", "descricao_predial", "area_medida_m2", "area_contratada_ha")):
            messages.append("artigo matricial extraído sem freguesia ou codigo_interno")
        else:
            messages.append("sem terreno; relação contrato-terreno não criada")
        if owner_id and parcel_id:
            right, problem = _right_type(entry.get("tipo_direito_terreno"))
            if problem:
                messages.append(problem)
            _upsert(workbook["db.ProprietarioTerreno"], TABLE_SCHEMAS["db.ProprietarioTerreno"], ("relacao_id",), {
                "relacao_id": stable_id("REL", "owner-parcel", owner_id, parcel_id), "proprietario_id": owner_id,
                "terreno_id": parcel_id, "tipo_direito": right, "quota": text(entry.get("quota")),
                "origem": "STEP4.0_ENTRADA_RAPIDA", "observacoes": entry.get("observacoes"),
            })
        if owner_id and parcel_id and not messages:
            return PROCESSING_DONE, "processada", ids
        return PROCESSING_PARTIAL, " | ".join(messages or ["associação proprietário-terreno não confirmada"]), ids

    def _owner(self, workbook, entry: Mapping[str, Any], reference: str) -> str:
        name = text(entry.get("nome_proprietario"))
        if not name:
            return ""
        nif = normalise_nif(entry.get("nif"))
        document = text(entry.get("documento_id"))
        owner_id = stable_id("PROP", "nif", nif) if valid_nif(nif) else (stable_id("PROP", "document", document) if document else stable_id("PROP", "name-contract", name, reference))
        owner_type, _ = _owner_type(entry.get("tipo_proprietario"))
        _upsert(workbook["db.Proprietario"], TABLE_SCHEMAS["db.Proprietario"], ("proprietario_id",), {
            "proprietario_id": owner_id, "nome_canonico": name, "tipo": owner_type,
            "documento_id": document, "nif": nif if valid_nif(nif) else "", "morada": entry.get("morada"), "estado": "ATIVO",
            "origem": "STEP4.0_ENTRADA_RAPIDA", "observacoes": entry.get("observacoes"),
        })
        return owner_id

    def _contacts_and_bank(self, workbook, entry: Mapping[str, Any], owner_id: str, messages: list[str]) -> None:
        for contact_type, field in (("TELEFONE", "telefone"), ("EMAIL", "email")):
            value = text(entry.get(field))
            if value:
                _upsert(workbook["db.Contacto"], TABLE_SCHEMAS["db.Contacto"], ("contacto_id",), {
                    "contacto_id": stable_id("CONT", owner_id, contact_type, value), "proprietario_id": owner_id,
                    "tipo": contact_type, "valor": value, "principal": True, "estado": "ATIVO",
                    "origem": "STEP4.0_ENTRADA_RAPIDA", "observacoes": "",
                })
        iban = normalise_iban(entry.get("iban"))
        if iban and valid_iban(iban):
            _upsert(workbook["db.ContaBancaria"], TABLE_SCHEMAS["db.ContaBancaria"], ("conta_id",), {
                "conta_id": stable_id("BANK", owner_id, iban), "proprietario_id": owner_id, "iban": iban,
                "origem": "STEP4.0_ENTRADA_RAPIDA", "observacoes": "",
            })
        elif iban:
            messages.append("IBAN inválido; não publicado")

    def _parcel(self, workbook, entry: Mapping[str, Any], source_row: int) -> str:
        code, parish, article, part = (text(entry.get(name)) for name in ("codigo_interno", "freguesia", "artigo_matricial", "parte"))
        if not code and not (parish and article):
            return ""
        parcel_id = stable_id("PAR", code, parish, article, part)
        at_m2 = as_decimal(entry.get("area_at_m2"))
        measured_m2 = as_decimal(entry.get("area_medida_m2"))
        documentary_ha = as_decimal(entry.get("area_contratada_ha"))
        at_ha = at_m2 / Decimal("10000") if at_m2 is not None else None
        measured_ha = measured_m2 / Decimal("10000") if measured_m2 is not None else None
        if documentary_ha is not None:
            considered, source, validation = documentary_ha, "AREA_DOCUMENTAL_PARCELA", "PENDENTE_VALIDACAO_DOCUMENTAL"
        elif at_ha is not None:
            considered, source, validation = at_ha, "AREA_AT_CONVERTIDA_ORIGINAL", "PENDENTE_VALIDACAO_DOCUMENTAL"
        else:
            considered, source, validation = None, "SEM_AREA_CONSIDERADA", "SEM_AREA_DOCUMENTAL"
        group_id = stable_id("TER", code or parish, parish, article)
        _upsert(workbook["db.Terreno"], TABLE_SCHEMAS["db.Terreno"], ("parcela_id",), {
            "parcela_id": parcel_id, "terreno_grupo_id": group_id, "codigo_interno": code,
            "distrito": entry.get("distrito"), "concelho": entry.get("concelho"), "freguesia": parish,
            "local": entry.get("local"), "artigo_matricial": article, "parte": part,
            "area_at_m2": at_m2, "area_at_ha": at_ha, "descricao_predial": entry.get("descricao_predial"),
            "area_medida_m2": measured_m2, "area_medida_ha": measured_ha,
            "area_registada_original_ha": at_ha, "area_documental_parcela_ha": documentary_ha,
            "area_considerada_ha": considered, "fonte_area_considerada": source,
            "status_validacao_area": validation, "observacoes": entry.get("observacoes"),
            "origem": "STEP4.0_ENTRADA_RAPIDA", "linha_origem": source_row,
        })
        return parcel_id

    def _reconcile_contract_areas(self, workbook) -> None:
        terrain_by_id = {row.get("parcela_id"): row for row in _rows(workbook["db.Terreno"])}
        relations: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for relation in _rows(workbook["db.ContratoTerreno"]):
            if text(relation.get("contrato_id")):
                relations[text(relation["contrato_id"])].append(relation)
        for contract in _rows(workbook["db.Contrato"]):
            contract_id = text(contract.get("contrato_id"))
            parcel_ids = {text(item.get("parcela_id")) for item in relations[contract_id]}
            values = [as_decimal(terrain_by_id[item].get("area_considerada_ha")) for item in parcel_ids if item in terrain_by_id]
            values = [item for item in values if item is not None]
            if values:
                total = sum(values, Decimal("0"))
                _upsert(workbook["db.Contrato"], TABLE_SCHEMAS["db.Contrato"], ("contrato_id",), {
                    "contrato_id": contract_id, "area_parcelas_considerada_ha": total,
                    "status_area_documental": "SOMA_PARCELAS_DISPONIVEL",
                })


def _authority_for(category: object, field: str) -> int:
    category = text(category).lower()
    order = {
        "property": {"property_document": 100, "lease_contract": 70, "bank_details": 20},
        "owner": {"property_document": 100, "lease_contract": 80, "bank_details": 30},
        "iban": {"bank_details": 100, "lease_contract": 80, "payment_proof": 70},
    }
    return order.get(field, {}).get(category, 10)


def _table_name(sheet_name: str) -> str:
    return "tbl_" + re.sub(r"[^A-Za-z0-9_]", "_", sheet_name.replace("audit.", "audit_"))


def _headers(sheet) -> dict[str, int]:
    return {text(cell.value): index for index, cell in enumerate(sheet[1], start=1) if text(cell.value)}


def _ensure_sheet_columns(sheet, columns: tuple[str, ...], table_name: str) -> None:
    existing_headers = [text(cell.value) for cell in sheet[1]] if sheet.max_row else []
    if tuple(existing_headers[:len(columns)]) != columns or sheet.max_column != len(columns):
        old_rows = _rows(sheet)
        styles = [copy(cell._style) for cell in sheet[1][:min(sheet.max_column, len(columns))]] if sheet.max_row else []
        if sheet.max_row:
            sheet.delete_rows(1, sheet.max_row)
        if sheet.max_column:
            sheet.delete_cols(1, sheet.max_column)
        sheet.append(columns)
        for index, style in enumerate(styles, start=1):
            sheet.cell(1, index)._style = style
        legacy_map = LEGACY_COLUMN_MAP.get(sheet.title, {})
        for row in old_rows:
            migrated = dict(row)
            for old, new in legacy_map.items():
                if not text(migrated.get(new)) and text(migrated.get(old)):
                    migrated[new] = migrated[old]
            sheet.append([_excel_value(migrated.get(column, ""), column) for column in columns])
    sheet.freeze_panes = "A2"
    _ensure_table(sheet, table_name, len(columns))
    for index, column in enumerate(columns, start=1):
        if column in TEXT_COLUMNS:
            sheet.column_dimensions[_letter(index)].width = max(sheet.column_dimensions[_letter(index)].width or 0, 16)
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, index).number_format = "@"
        elif column in DATE_COLUMNS:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, index).number_format = "yyyy-mm-dd"
        elif column in DATETIME_COLUMNS:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, index).number_format = "yyyy-mm-dd hh:mm:ss"


def _ensure_table(sheet, name: str, width: int) -> None:
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    if name in sheet.tables:
        table = sheet.tables[name]
    elif sheet.tables:
        table = next(iter(sheet.tables.values()))
    else:
        table = Table(displayName=name, ref=f"A1:{_letter(width)}{max(1, sheet.max_row)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    table.ref = f"A1:{_letter(width)}{max(1, sheet.max_row)}"
    if table.autoFilter is None:
        table.autoFilter = AutoFilter(ref=table.ref)
    else:
        table.autoFilter.ref = table.ref


def _rows(sheet) -> list[dict[str, Any]]:
    headers = _headers(sheet)
    return [{name: sheet.cell(row, column).value for name, column in headers.items()} for row in range(2, sheet.max_row + 1) if any(text(sheet.cell(row, column).value) for column in headers.values())]


def _row_by_key(sheet, columns: tuple[str, ...], keys: tuple[str, ...], values: tuple[object, ...]) -> dict[str, Any] | None:
    for row in _rows(sheet):
        if all(comparison_key(row.get(key)) == comparison_key(value) for key, value in zip(keys, values)):
            return row
    return None


def _upsert(sheet, columns: tuple[str, ...], keys: tuple[str, ...], values: Mapping[str, Any], *, replace: bool = False) -> tuple[int, bool]:
    headers = _headers(sheet)
    key_values = tuple(values.get(key, "") for key in keys)
    row_number = None
    for candidate in range(2, sheet.max_row + 1):
        if all(comparison_key(sheet.cell(candidate, headers[key]).value) == comparison_key(value) for key, value in zip(keys, key_values)):
            row_number = candidate
            break
    created = row_number is None
    if created:
        sheet.append([_excel_value(values.get(column, ""), column) for column in columns])
        row_number = sheet.max_row
    else:
        for column, value in values.items():
            if column not in headers or value in (None, ""):
                continue
            cell = sheet.cell(row_number, headers[column])
            if replace or not text(cell.value):
                cell.value = _excel_value(value, column)
    for column, index in headers.items():
        if column in TEXT_COLUMNS:
            sheet.cell(row_number, index).number_format = "@"
        elif column in DATE_COLUMNS:
            sheet.cell(row_number, index).number_format = "yyyy-mm-dd"
        elif column in DATETIME_COLUMNS:
            sheet.cell(row_number, index).number_format = "yyyy-mm-dd hh:mm:ss"
    _ensure_table(sheet, _table_name(sheet.title), len(columns))
    return row_number, created


def _as_date(value: object) -> date | str:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    candidate = text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(candidate, pattern).date()
        except ValueError:
            pass
    return candidate


def _excel_value(value: object, column: str = "") -> object:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    if column in TEXT_COLUMNS and value is not None:
        return str(value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


__all__ = [
    "AUDIT_SCHEMAS", "ENTRY_COLUMNS", "NORMALIZED_WORKBOOK_SCHEMA_VERSION", "NormalizedWorkbookRegister",
    "PROCESSING_DONE", "PROCESSING_PARTIAL", "PROCESSING_PENDING", "PROCESSING_REVIEW", "TABLE_SCHEMAS",
    "build_staging_rows", "normalise_iban", "normalise_nif", "stable_id", "valid_iban", "valid_nif",
]
