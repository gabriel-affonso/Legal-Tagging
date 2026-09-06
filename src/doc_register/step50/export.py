"""Compatibility projection. Export into a fresh workbook, never merge stale facts."""
from collections import defaultdict
from datetime import datetime, timezone
import json
from pathlib import Path
import tempfile

from ..models import ExtractionResult, PdfCandidate
from ..registry import ExcelRegister
from ..step40 import NormalizedWorkbookRegister
from .storage import locked


def project(result):
    types = {s['source_type'] for s in result.get('segments', [])}
    category = 'lease_contract' if types & {'contract', 'amendment'} else 'property_document' if types & {'caderneta', 'registry'} else 'unknown'
    out = ExtractionResult(document_category=category, needs_review='Yes', human_review_required='Yes',
                           processing_status='STEP5_REVIEW', text_source='step5_canonical',
                           extraction_notes='Step 5.0; only accepted facts projected. Full evidence and exact roles in audit.Step50. Whole-document quality uncalibrated.')
    groups = defaultdict(dict)
    for field in result['fields']:
        if field['status'] == 'accepted':
            groups[(field['logical_document'], field['entity'])][field['field']] = field['value']
    # Scalar legacy columns cannot truthfully collapse multiple contracts/periods.
    contracts = [values for (logical, entity), values in groups.items() if entity == 'contract' and logical.startswith('contract-')]
    if len(contracts) == 1:
        for key in ['signed_date', 'contract_start_date', 'contract_end_date', 'contract_type']:
            setattr(out, key, contracts[0].get(key, ''))
        rent = contracts[0].get('rent')
        if rent:
            out.rent_amount, out.rent_currency = rent['amount'], rent['currency']
            out.rent_frequency, out.rent_basis, out.rent_condition = rent['frequency'], rent['basis'], rent['condition']
    properties = []
    lessors, lessees = [], []
    for (logical, entity), values in groups.items():
        if values.get('name'):
            owner = {'name': values['name'], 'nif': values.get('tax_id', ''), 'role': 'OUTRO'}
            if values.get('role') == 'lessor':
                lessors.append(owner)
            elif values.get('role') == 'lessee':
                lessees.append(owner)
        if values.get('property_article'):
            prop = {key: value for key, value in values.items() if key.startswith('property_') and not isinstance(value, dict)}
            area = values.get('property_total_area')
            if area:
                prop['property_total_area'] = area['amount'] + ' ' + area['unit']
            prop['source'] = logical + ':' + entity
            properties.append(prop)
    out.contract_lessor, out.contract_lessee = lessors, lessees
    out.lessor, out.lessee = '; '.join(p['name'] for p in lessors), '; '.join(p['name'] for p in lessees)
    if len(lessors) == 1:
        out.lessor_tax_id = lessors[0]['nif']
    if len(lessees) == 1:
        out.lessee_tax_id = lessees[0]['nif']
    out.raw_json = {'final_contract_resolution': {'properties': properties}, 'step50_run_id': result['run_id']}
    return out


def export_results(paths, destination: Path, output_format='both', reference=''):
    from openpyxl import load_workbook
    destination = destination.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with locked(destination.with_suffix('.lock')):
        if destination.exists():
            raise FileExistsError('Use a new export path; operational/human workbooks are never overwritten.')
        with tempfile.TemporaryDirectory(dir=destination.parent, prefix='.step50-export-') as temporary:
            target = Path(temporary) / 'export.xlsx'
            for path in paths:
                result = json.loads(Path(path).read_text())
                now = datetime.now(timezone.utc)
                original = Path(result['document']['original_path'])
                candidate = PdfCandidate(original, original, result['document']['sha256'], now, now)
                extraction = project(result)
                if output_format in {'normalized', 'both'}:
                    NormalizedWorkbookRegister(target).publish(candidate, extraction, reference_context=reference)
                if output_format in {'legacy', 'both'}:
                    ExcelRegister(target).append(candidate, extraction)
                workbook = load_workbook(target)
                sheet = workbook['audit.Step50'] if 'audit.Step50' in workbook else workbook.create_sheet('audit.Step50')
                if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                    sheet.append(['sha256', 'run_id', 'logical_document', 'entity', 'field', 'status', 'value', 'acceptance', 'candidate_ids', 'result_path'])
                for f in result['fields']:
                    sheet.append([result['document']['sha256'], result['run_id'], f['logical_document'], f['entity'], f['field'], f['status'],
                                  json.dumps(f['value'], ensure_ascii=False), f['acceptance'], ','.join(f['candidate_ids']), str(Path(path).resolve())])
                # Treat extracted strings as data, including formula-looking entity names.
                for ws in workbook:
                    for row in ws:
                        for cell in row:
                            if cell.data_type == 'f':
                                cell.data_type = 's'
                workbook.save(target)
                workbook.close()
            if not target.exists():
                raise ValueError('no_results_to_export')
            target.replace(destination)
    return destination
